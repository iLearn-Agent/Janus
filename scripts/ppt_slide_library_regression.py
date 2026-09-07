from __future__ import annotations

import base64
import io
import json
import os
import re
import tempfile
import threading
import time
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
ASSET_ROOT = ROOT / "assets"
SERVICE = ROOT / "src" / "main" / "ppt_service"
TMP = Path(tempfile.gettempdir()) / "janus-ppt-regression"

import sys

sys.path.insert(0, str(SERVICE))

import ppt_debug_service as ppt_service  # noqa: E402
import ppt_pipeline.office_conversion as office_conversion  # noqa: E402
from slide_library_renderer import render_slide_library_presentation, validate_slide_library  # noqa: E402
from ppt_debug_service import (  # noqa: E402
    SlideSpec,
    _adapt_sparse_specs_for_quality_repair,
    _map_source_visuals_to_slides,
    _ppt_image_binding_issues,
    _ppt_text_quality_issues,
    _selected_image_slide_indices,
    resolve_ppt_template,
)


def spec(title: str, message: str = "") -> SimpleNamespace:
    return SimpleNamespace(title=title, message=message, visual="", speaker_note="")


def named(slide, prefix: str):
    return [shape for shape in slide.shapes if str(shape.name).startswith(prefix)]


def visible_text(slide) -> str:
    return " ".join(str(getattr(shape, "text", "") or "") for shape in slide.shapes)


def min_font(shape) -> float:
    sizes = [
        float(run.font.size.pt)
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.font.size is not None
    ]
    return min(sizes) if sizes else 0.0


def run_typefaces(shape) -> list[tuple[str | None, str | None]]:
    result: list[tuple[str | None, str | None]] = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            properties = run._r.get_or_add_rPr()
            east_asian = properties.find(qn("a:ea"))
            result.append((run.font.name, east_asian.get("typeface") if east_asian is not None else None))
    return result


def has_soft_shadow(shape) -> bool:
    effects = shape._element.spPr.find(qn("a:effectLst"))
    return effects is not None and effects.find(qn("a:outerShdw")) is not None


def connector_endpoints(shape):
    transform = shape._element.spPr.xfrm
    left = float(shape.left)
    top = float(shape.top)
    right = float(shape.left + shape.width)
    bottom = float(shape.top + shape.height)
    start_x, end_x = (right, left) if transform.get("flipH") in {"1", "true"} else (left, right)
    start_y, end_y = (bottom, top) if transform.get("flipV") in {"1", "true"} else (top, bottom)
    return (start_x, start_y), (end_x, end_y)


def point_on_shape_border(point, shape, tolerance=3.0):
    x, y = point
    left = float(shape.left)
    top = float(shape.top)
    right = float(shape.left + shape.width)
    bottom = float(shape.top + shape.height)
    within_x = left - tolerance <= x <= right + tolerance
    within_y = top - tolerance <= y <= bottom + tolerance
    return within_x and within_y and (
        abs(x - left) <= tolerance
        or abs(x - right) <= tolerance
        or abs(y - top) <= tolerance
        or abs(y - bottom) <= tolerance
    )


def semantic_bottom(slide) -> float:
    prefixes = (
        "basic-", "motivation-", "challenge-", "pipeline-", "loop-", "bench-", "big-", "bars-", "bar",
        "leader-", "ablation-", "evidence-", "case-", "summary-", "target-", "domain-", "route-", "wp-",
        "eval-", "risk-", "roadmap-", "media-", "janus-",
    )
    bottoms = [
        float(shape.top + shape.height) / 914400
        for shape in slide.shapes
        if str(shape.name).startswith(prefixes)
    ]
    return max(bottoms, default=0.0)


def run(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    source_prs = Presentation(str(template))
    for index, slide in enumerate(list(source_prs.slides)[1:], start=2):
        bottom = semantic_bottom(slide)
        if bottom:
            assert bottom >= 6.45, f"template slide {index} does not use the expanded lower canvas"
    specs = [
        spec("Regression Cover", "Semantic page-library regression"),
        spec("Two challenges should use two cards"),
        spec("Three steps should use three process cards"),
        spec("Three stages should populate the loop"),
        spec("Metric groups should populate KPI cards"),
        spec("Two evidence cards should remove four slots"),
        spec("Three cases should remove two unused slots"),
        spec("Media without an image should become a pipeline"),
        spec("Two takeaways should use two cards"),
    ]
    layout_ids = [
        "cover",
        "challenge_map",
        "method_pipeline",
        "method_loop",
        "benchmark_metrics",
        "evidence_grid",
        "case_gallery",
        "media_showcase",
        "summary_takeaways",
    ]
    payloads = [
        {"title": specs[0].title},
        {
            "title": specs[1].title,
            "nodes": [
                {"label": "Grounding", "detail": "Keep outputs catalog-valid"},
                {"label": "Latency", "detail": "Keep decoding responsive"},
            ],
            "objective": "Controlled generation",
        },
        {
            "title": specs[2].title,
            "input": ["Context", "Catalog"],
            "steps": ["Encode", "Generate", "Validate"],
            "output": ["Recommendation"],
        },
        {
            "title": specs[3].title,
            "stages": [
                {"label": "Train", "detail": "Learn sequences"},
                {"label": "Serve", "detail": "Generate safely"},
                {"label": "Observe", "detail": "Collect feedback"},
            ],
            "objective": "Improve user utility",
        },
        {
            "title": specs[4].title,
            "kpis": [
                {"group": "Ranking", "metrics": ["Recall", "NDCG"]},
                {"group": "System", "metrics": ["Latency", "Cost"]},
            ],
            "protocol": "Use compatible data splits",
        },
        {
            "title": specs[5].title,
            "cards": [
                {"title": "ID tokens", "points": ["Compact", "Catalog-specific"]},
                {"title": "Semantic codes", "points": ["Structured", "Transferable"]},
            ],
        },
        {
            "title": specs[6].title,
            "cards": [
                {"title": "Sequential", "points": ["Predict the next item"]},
                {"title": "Conversational", "points": ["Clarify preferences"]},
                {"title": "Multi-task", "points": ["Recommend and explain"]},
            ],
        },
        {
            "title": specs[7].title,
            "input": "User request",
            "steps": ["Interpret", "Ground", "Generate"],
            "output": "Valid response",
        },
        {
            "title": specs[8].title,
            "points": ["Ground outputs", "Evaluate utility"],
            "conclusion": "Use generation where it adds value",
        },
    ]

    progress_events: list[dict[str, object]] = []
    prs, warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=layout_ids,
        content_payloads=payloads,
        image_paths={},
        on_progress=progress_events.append,
    )
    assert [event["current_slide"] for event in progress_events] == list(range(1, len(specs) + 1))
    assert all(event["total_slides"] == len(specs) for event in progress_events)
    assert [event["phase_current"] for event in progress_events] == list(range(1, len(specs) + 1))
    assert all(event["phase_total"] == len(specs) for event in progress_events)
    assert len(prs.slides) == len(specs)
    assert len(named(prs.slides[1], "challenge-card")) == 2
    assert not named(prs.slides[1], "challenge-note")
    challenge_cards = named(prs.slides[1], "challenge-card")
    challenge_core = named(prs.slides[1], "challenge-core")[0]
    for link, card in zip(named(prs.slides[1], "challenge-link"), challenge_cards):
        card_endpoint, core_endpoint = connector_endpoints(link)
        assert point_on_shape_border(card_endpoint, card)
        assert point_on_shape_border(core_endpoint, challenge_core)
    assert len(named(prs.slides[2], "pipeline-step-")) == 3
    assert not named(prs.slides[2], "pipeline-grid-"), "pipeline must not repeat step labels in an automatic support table"
    assert named(prs.slides[2], "pipeline-step-")[0].height / 914400 >= 4.4
    assert len(named(prs.slides[3], "loop-node-")) == 3
    assert len(named(prs.slides[4], "bench-kpi1-value")) == 1
    assert len(named(prs.slides[5], "evidence-img-")) == 2
    assert len(named(prs.slides[6], "case-thumb-")) == 4  # two cards plus their caption shapes
    assert named(prs.slides[7], "tpl-title-17"), "repeated media fallback should use the technical_route alternate"
    assert len(named(prs.slides[8], "summary-card")) == 2
    summary_cards = named(prs.slides[8], "summary-card")
    assert summary_cards[0].width / 914400 >= 5.4
    assert summary_cards[1].width / 914400 >= 5.4
    # Supporting cards stay flat and quiet; only the primary visual anchor
    # receives a restrained depth cue.
    assert not has_soft_shadow(named(prs.slides[1], "challenge-card")[0])
    assert has_soft_shadow(named(prs.slides[1], "challenge-core")[0])
    assert not has_soft_shadow(named(prs.slides[8], "summary-card")[0])
    assert has_soft_shadow(named(prs.slides[8], "summary-next")[0])
    assert any("media_showcase" in warning and "method_pipeline" in warning for warning in warnings)
    assert min_font(named(prs.slides[2], "pipeline-step-")[0]) >= 14

    placeholder = re.compile(r"\[[^\]]+\]|placeholder|(?:证据|案例)\s*\d+", re.IGNORECASE)
    for index, slide in enumerate(list(prs.slides)[1:], start=2):
        assert not placeholder.search(visible_text(slide)), f"slide {index} contains placeholder text"

    out = TMP / f"regression-{template.stem}.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    reopened = Presentation(out)
    assert len(reopened.slides) == len(specs)
    out.unlink(missing_ok=True)


def font_and_image_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    specs = [spec("Readable Cover", "Readable body typography"), spec("Readable comparison title")]
    payloads = [
        {"title": specs[0].title},
        {
            "title": specs[1].title,
            "current": {"label": "Current", "points": ["Score a fixed candidate set", "Optimize ranking"]},
            "target": {"label": "Target", "points": ["Generate grounded sequences", "Support interaction"]},
            "gap": "Generation needs stronger grounding and control",
        },
    ]
    prs, _warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=["cover", "motivation_compare"],
        content_payloads=payloads,
        image_paths={},
    )
    body = prs.slides[1]
    if "哈工深" in template_name:
        subtitle = next(
            shape for shape in prs.slides[0].shapes
            if getattr(shape, "has_text_frame", False) and shape.text == "Readable body typography"
        )
        subtitle_colors = {
            str(run.font.color.rgb)
            for paragraph in subtitle.text_frame.paragraphs
            for run in paragraph.runs
            if run.text.strip()
        }
        assert subtitle_colors == {"0F3A56"}, "HITSZ cover subtitle must remain visible on the white cover region"
    kicker = named(body, "tpl-kicker-label")[0]
    assert kicker.text == "WHY NOW?"
    assert min_font(kicker) >= 14
    assert kicker.fill.type == MSO_FILL_TYPE.BACKGROUND
    assert kicker.line.fill.type == MSO_FILL_TYPE.BACKGROUND
    assert kicker.top + kicker.height < min(shape.top for shape in named(body, "motivation-current"))
    assert not named(body, "motivation-mini-table-"), "optional duplicate mini table must be removed"
    assert min_font(named(body, "tpl-title-")[0]) >= 24
    assert min_font(named(body, "motivation-current")[0]) >= 16
    assert all(latin == "Comic Sans MS" and east_asian == "Microsoft YaHei" for latin, east_asian in run_typefaces(named(body, "motivation-current")[0]))
    assert named(body, "motivation-gap")[0].width / 914400 >= 2.8
    assert named(body, "motivation-current")[0].text_frame.vertical_anchor == MSO_ANCHOR.TOP

    image_path = TMP / "ppt-regression-image.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (640, 360), "#2F6B8A").save(image_path)
    with_image, warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=["cover", "motivation_compare"],
        content_payloads=payloads,
        image_paths={2: image_path},
    )
    assert named(with_image.slides[1], "tpl-title-2"), "non-image layout must switch to basic_content"
    assert any(shape.name == "janus-basic-image-image" for shape in with_image.slides[1].shapes)
    assert any("没有图片槽位" in warning and "basic_content" in warning for warning in warnings)

    repeated_basic, repeated_warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=[spec("Cover"), spec("Visual one"), spec("Visual two")],
        layout_ids=["cover", "basic_content", "basic_content"],
        content_payloads=[{}, {"points": ["Point A", "Point B"]}, {"points": ["Point C", "Point D"]}],
        image_paths={2: image_path, 3: image_path},
    )
    first_basic = repeated_basic.slides[1]
    mirrored_basic = repeated_basic.slides[2]
    assert not named(first_basic, "tpl-kicker-label")
    assert not named(first_basic, "basic-caption")
    assert named(first_basic, "janus-basic-image-image")[0].width > named(first_basic, "basic-text-card")[0].width
    assert named(first_basic, "janus-basic-image-image")[0].width >= named(first_basic, "basic-text-card")[0].width * 2
    assert named(mirrored_basic, "tpl-title-basic-mirror")
    assert any("basic_content_mirror" in warning for warning in repeated_warnings)

    for image_layout in ("evidence_grid", "case_gallery"):
        single_image, single_warnings = render_slide_library_presentation(
            template_path=template,
            template_id="scut" if "华工" in template_name else "hitsz",
            specs=specs,
            layout_ids=["cover", image_layout],
            content_payloads=payloads,
            image_paths={2: image_path},
        )
        assert named(single_image.slides[1], "tpl-title-2"), "single-image gallery must use basic_content"
        assert any(shape.name == "janus-basic-image-image" for shape in single_image.slides[1].shapes)
        assert not named(single_image.slides[1], "case-thumb-"), "single-image layout must not retain empty gallery slots"
        assert any("单图版式" in warning for warning in single_warnings)

    multi_paths: list[Path] = []
    for index, color in enumerate(("#1D4ED8", "#0F766E", "#D97706", "#B91C1C"), start=1):
        path = TMP / f"multi-figure-{index}.png"
        Image.new("RGB", (640, 360), color).save(path)
        multi_paths.append(path)
    multi_specs = [spec("Cover"), spec("Four evidence figures")]
    evidence, evidence_warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=multi_specs,
        layout_ids=["cover", "evidence_grid"],
        content_payloads=[{}, {"captions": ["Figure A", "Figure B", "Figure C", "Figure D"]}],
        image_paths={2: multi_paths},
    )
    assert named(evidence.slides[1], "tpl-title-12"), "multi-image evidence must keep the gallery layout"
    assert len(named(evidence.slides[1], "janus-evidence-img-")) == 4
    assert not any("单图版式" in warning for warning in evidence_warnings)

    cases, _case_warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=[spec("Cover"), spec("Three cases")],
        layout_ids=["cover", "case_gallery"],
        content_payloads=[{}, {"captions": ["Main", "Case B", "Case C"]}],
        image_paths={2: multi_paths[:3]},
    )
    assert named(cases.slides[1], "tpl-title-13")
    assert len([shape for shape in cases.slides[1].shapes if shape.name.startswith("janus-") and shape.name.endswith("-image")]) == 3
    for path in multi_paths:
        path.unlink(missing_ok=True)
    image_path.unlink(missing_ok=True)


def image_caption_and_density_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    template_id = "scut" if "华工" in template_name else "hitsz"
    diagram_path = TMP / "contain-diagram.png"
    photo_path = TMP / "crop-photo.png"
    extra_photo_path = TMP / "crop-photo-extra.png"
    diagram_path.parent.mkdir(parents=True, exist_ok=True)

    diagram = Image.new("RGB", (600, 800), "white")
    diagram_draw = ImageDraw.Draw(diagram)
    for row in range(9):
        y = 45 + row * 78
        diagram_draw.rectangle((55, y, 545, y + 46), outline=(45, 92, 132), width=4)
        diagram_draw.line((80, y + 17, 420, y + 17), fill=(62, 112, 151), width=3)
        diagram_draw.line((80, y + 31, 500, y + 31), fill=(118, 145, 166), width=2)
    diagram.save(diagram_path)

    photo = Image.new("RGB", (600, 800))
    photo_draw = ImageDraw.Draw(photo)
    for y in range(800):
        photo_draw.line((0, y, 599, y), fill=(25 + y % 170, 55 + (y * 2) % 150, 115 + (y * 3) % 120))
    photo.save(photo_path)
    photo.transpose(Image.Transpose.FLIP_LEFT_RIGHT).save(extra_photo_path)

    density_specs = [spec("Cover"), spec("Sparse statement"), spec("Dense explanation")]
    density, _ = render_slide_library_presentation(
        template_path=template,
        template_id=template_id,
        specs=density_specs,
        layout_ids=["cover", "basic_content", "basic_content"],
        content_payloads=[
            {},
            {"conclusion": "One decision should remain memorable", "points": ["Use evidence to support the decision"]},
            {
                "points": [
                    "Define the problem boundary before selecting a method",
                    "Separate assumptions from directly observed evidence",
                    "Compare alternatives under the same evaluation protocol",
                    "Record failure cases and operational constraints",
                    "Close with a concrete decision and next action",
                ]
            },
        ],
        image_paths={},
    )
    sparse = named(density.slides[1], "basic-text-card")[0]
    dense = named(density.slides[2], "basic-text-card")[0]
    assert sparse.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
    assert min_font(sparse) >= 20
    assert all(run.font.bold for run in sparse.text_frame.paragraphs[0].runs)
    assert dense.text_frame.vertical_anchor == MSO_ANCHOR.TOP
    assert min_font(dense) >= 15

    evidence_specs = [spec("Cover"), spec("Figures with conclusions")]
    evidence, _ = render_slide_library_presentation(
        template_path=template,
        template_id=template_id,
        specs=evidence_specs,
        layout_ids=["cover", "evidence_grid"],
        content_payloads=[
            {},
            {
                "captions": [
                    {"conclusion": "The full evaluation matrix must remain readable", "source": "Offline experiment log"},
                    {"caption": "The field photo provides contextual evidence", "source": "Project team"},
                ]
            },
        ],
        image_paths={2: [diagram_path, photo_path]},
    )
    evidence_slide = evidence.slides[1]
    contained = named(evidence_slide, "janus-evidence-img-0-0-image")[0]
    cropped = named(evidence_slide, "janus-evidence-img-0-1-image")[0]
    assert abs(contained.width / contained.height - diagram.width / diagram.height) < 0.01
    assert not any((contained.crop_left, contained.crop_right, contained.crop_top, contained.crop_bottom))
    assert cropped.crop_top > 0 or cropped.crop_bottom > 0
    first_caption = named(evidence_slide, "evidence-caption-0-0")[0]
    assert first_caption.text.startswith("Figure 1")
    assert "Source: Offline experiment log" in first_caption.text
    assert first_caption.fill.type == MSO_FILL_TYPE.BACKGROUND
    assert first_caption.line.fill.type == MSO_FILL_TYPE.BACKGROUND
    assert first_caption.text_frame.paragraphs[0].runs[0].font.bold
    assert first_caption.text_frame.paragraphs[-1].runs[0].font.size.pt < first_caption.text_frame.paragraphs[0].runs[0].font.size.pt

    cases, _ = render_slide_library_presentation(
        template_path=template,
        template_id=template_id,
        specs=[spec("Cover"), spec("Primary and supporting cases")],
        layout_ids=["cover", "case_gallery"],
        content_payloads=[{}, {"captions": [{"caption": "Primary case", "source": "Project archive"}]}],
        image_paths={2: [photo_path, extra_photo_path, photo_path]},
    )
    case_slide = cases.slides[1]
    assert "Source: Project archive" in named(case_slide, "case-caption-main")[0].text
    assert not named(case_slide, "case-thumb-caption-")
    assert all(shape.height / 914400 >= 2.4 for shape in named(case_slide, "janus-case-thumb-"))

    media, _ = render_slide_library_presentation(
        template_path=template,
        template_id=template_id,
        specs=[spec("Cover"), spec("Full-width proof object")],
        layout_ids=["cover", "media_showcase"],
        content_payloads=[{}, {"proof_object": "System architecture overview", "source": "Design specification"}],
        image_paths={2: diagram_path},
    )
    media_slide = media.slides[1]
    assert named(media_slide, "media-caption")
    assert "Source: Design specification" in named(media_slide, "media-caption")[0].text
    media_picture = named(media_slide, "janus-media-showcase-main-image")[0]
    assert not any((media_picture.crop_left, media_picture.crop_right, media_picture.crop_top, media_picture.crop_bottom))

    for path in (diagram_path, photo_path, extra_photo_path):
        path.unlink(missing_ok=True)


def data_visual_hierarchy_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    specs = [spec("Cover"), spec("Benchmark"), spec("Results"), spec("Leaderboard"), spec("Evaluation")]
    payloads = [
        {},
        {
            "kpis": [
                {"value": "94.8%", "label": "Accuracy", "note": "Representative test split"},
                {"value": "1.8x", "label": "Throughput", "note": "Same serving budget"},
                {"value": "-23%", "label": "Cost", "note": "End-to-end estimate"},
            ],
            "protocol": "Report the protocol once; do not repeat KPI labels in a synthetic table.",
        },
        {
            "chart_title": "Quality comparison",
            "bars": [
                {"label": "Baseline", "value": "61%"},
                {"label": "Ours", "value": "82%"},
                {"label": "Target", "value": "90%"},
            ],
            "points": ["The largest gain comes from grounded validation", "The remaining gap is concentrated in long-tail cases"],
        },
        {
            "headers": ["Rank", "Method", "Score", "Cost", "Comment"],
            "rows": [
                ["1", "Ours", "92.1", "Medium", "Best quality-cost tradeoff"],
                ["2", "Baseline A", "86.4", "Low", "Fast but less accurate"],
                ["3", "Baseline B", "84.2", "High", "Strong quality, high cost"],
            ],
        },
        {
            "kpis": [
                {"value": "92%", "label": "Quality", "note": "Primary acceptance metric"},
                {"value": "180ms", "label": "Latency", "note": "P95 serving latency"},
                {"value": "0.8x", "label": "Cost", "note": "Relative serving cost"},
            ],
            "scope": "Evaluate quality, system health, and evidence readiness together.",
            "status_rows": [
                {"item": "Offline benchmark", "status": "Ready", "evidence": "Reproducible report"},
                {"item": "Online pilot", "status": "Pending", "evidence": "Guarded experiment"},
            ],
        },
    ]
    prs, _warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=["cover", "benchmark_metrics", "results_bars", "leaderboard_table", "evaluation_dashboard"],
        content_payloads=payloads,
        image_paths={},
    )
    benchmark = prs.slides[1]
    assert not named(benchmark, "bench-table-"), "KPI-only benchmark must not retain a synthetic table"
    assert named(benchmark, "bench-kpi1")[0].height / 914400 >= 3.6
    assert min_font(named(benchmark, "bench-kpi1-value")[0]) >= 26

    bars = prs.slides[2]
    assert not named(bars, "bar4-")
    assert not named(bars, "bars-note2") and not named(bars, "bars-note3")
    assert named(bars, "bars-note1")[0].height / 914400 >= 5.0
    assert min_font(named(bars, "bar1-label")[0]) >= 12

    leaderboard = prs.slides[3]
    assert not named(leaderboard, "leader-note")
    last_cell = named(leaderboard, "leaderboard-grid-r1c5")[0]
    assert (last_cell.left + last_cell.width) / 914400 >= 12.5
    assert min_font(named(leaderboard, "leaderboard-grid-r2c2")[0]) >= 11

    evaluation = prs.slides[4]
    assert not named(evaluation, "eval-bar")
    assert named(evaluation, "eval-kpi1")[0].height / 914400 >= 2.0
    assert named(evaluation, "eval-scope")[0].width / 914400 >= 3.7


def semantic_emphasis_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    specs = [
        spec("Cover"),
        spec("Why change"),
        spec("Headline outcomes"),
        spec("Comparison result"),
        spec("Evaluation readiness"),
        spec("Decision summary"),
    ]
    payloads = [
        {},
        {
            "current": {"label": "Fragmented workflow", "value": "42%", "points": ["Evidence is scattered", "Decisions are difficult to trace"]},
            "target": {"label": "Closed-loop workflow", "value": "68%", "points": ["Evidence supports each gate", "Decisions remain auditable"]},
            "gap": "A shared evidence and ownership model is still missing",
        },
        {
            "kpis": [
                {"value": "-18%", "label": "Cost", "note": "Same serving envelope"},
                {"value": "+21 pts", "label": "Quality gain", "note": "Primary acceptance result", "primary": True},
                {"value": "1.6x", "label": "Throughput", "note": "P95 production load"},
            ],
            "interpretation": "Quality improves materially without sacrificing the operating envelope",
        },
        {
            "bars": [
                {"label": "Baseline", "value": "61%"},
                {"label": "Ours", "value": "82%"},
                {"label": "Target", "value": "90%"},
            ]
        },
        {
            "kpis": [
                {"value": "92%", "label": "Quality", "note": "Representative test split"},
                {"value": "96%", "label": "Validity", "note": "Primary acceptance metric"},
                {"value": "180ms", "label": "Latency", "note": "P95 serving latency"},
            ],
            "scope": "Use quality, validity, latency, and readiness evidence together.",
            "status_rows": [
                {"item": "Offline benchmark", "status": "Ready", "evidence": "Reproducible report"},
                {"item": "Online pilot", "status": "Pending", "evidence": "Guarded experiment"},
                {"item": "Scale gate", "status": "At risk", "evidence": "Cost evidence incomplete"},
            ],
        },
        {
            "findings": ["The controlled method produces the strongest verified result", "Most remaining errors concentrate in sparse scenarios"],
            "limitations": "The online sample is still too small for a scale decision",
            "next_step": "Run the guarded pilot and assign an owner to each acceptance gate",
        },
    ]
    prs, _ = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=["cover", "motivation_compare", "result_big_numbers", "results_bars", "evaluation_dashboard", "summary_takeaways"],
        content_payloads=payloads,
        image_paths={},
    )

    motivation = prs.slides[1]
    current = named(motivation, "motivation-current")[0]
    target = named(motivation, "motivation-target")[0]
    gap = named(motivation, "motivation-gap")[0]
    assert str(current.fill.fore_color.rgb) == "F7FAFC"
    assert str(target.fill.fore_color.rgb) == "EFF8F9"
    assert str(gap.fill.fore_color.rgb) == "FFF7E4"
    arrow_start, arrow_end = connector_endpoints(named(motivation, "motivation-arrow")[0])
    assert point_on_shape_border(arrow_start, current)
    assert point_on_shape_border(arrow_end, target)

    big = prs.slides[2]
    assert named(big, "big-kpi2")[0].width / 914400 >= 5.2
    assert named(big, "big-kpi1")[0].left == named(big, "big-kpi3")[0].left
    big_accents = [str(named(big, f"big-kpi{index}-accent")[0].fill.fore_color.rgb) for index in range(1, 4)]
    assert big_accents == ["8FB0C3", "008A9A", "8FB0C3"]
    assert "KEY FINDING" in named(big, "big-interpretation")[0].text

    bars = prs.slides[3]
    assert str(named(bars, "bar2-fill")[0].fill.fore_color.rgb) == "008A9A"
    assert str(named(bars, "bar1-fill")[0].fill.fore_color.rgb) == "8FB0C3"
    assert named(bars, "bar2-fill")[0].height > named(bars, "bar1-fill")[0].height
    assert "+21 pts" in named(bars, "bars-note1")[0].text

    evaluation = prs.slides[4]
    eval_accents = [str(named(evaluation, f"eval-kpi{index}-accent")[0].fill.fore_color.rgb) for index in range(1, 4)]
    assert eval_accents == ["8FB0C3", "008A9A", "8FB0C3"]
    assert str(named(evaluation, "eval-status-grid-r2c2")[0].fill.fore_color.rgb) == "EAF7EF"
    assert str(named(evaluation, "eval-status-grid-r3c2")[0].fill.fore_color.rgb) == "F2F5F7"
    assert str(named(evaluation, "eval-status-grid-r4c2")[0].fill.fore_color.rgb) == "FFF7E4"

    summary = prs.slides[5]
    assert named(summary, "summary-card1")[0].text.startswith("PRIMARY FINDING")
    assert named(summary, "summary-card3")[0].text.startswith("LIMITATION")
    assert str(named(summary, "summary-card1")[0].fill.fore_color.rgb) == "EFF8F9"
    assert str(named(summary, "summary-card3")[0].fill.fore_color.rgb) == "FFF7E4"
    assert str(named(summary, "summary-next")[0].fill.fore_color.rgb) == "0F3A56"
    assert "Run the guarded pilot" in named(summary, "summary-next")[0].text


def matrix_visual_hierarchy_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    specs = [spec("Cover"), spec("消融结果说明完整方案最稳定"), spec("任务包围绕交付物组织"), spec("高影响风险需要优先处理")]
    payloads = [
        {},
        {
            "headers": ["方案", "模块 A", "模块 B", "模块 C", "得分", "变化"],
            "rows": [
                ["完整方案", "✓", "✓", "✓", "92.4", "—"],
                ["移除 A", "—", "✓", "✓", "84.1", "-8.3"],
                ["移除 B", "✓", "—", "✓", "87.5", "-4.9"],
                ["移除 C", "✓", "✓", "—", "89.0", "-3.4"],
            ],
            "conclusion": "完整方案在质量与稳定性之间取得最佳平衡。",
        },
        {
            "headers": ["任务包", "负责人", "核心任务", "交付物", "指标", "时间"],
            "rows": [
                ["WP1", "算法组", "建立基线", "评测报告", "可复现", "Q1"],
                ["WP2", "平台组", "实现服务", "原型系统", "P95<200ms", "Q2"],
                ["WP3", "测试组", "组织验收", "验收材料", "覆盖关键场景", "Q3"],
            ],
            "note": "当前优先推进 WP1，后续任务以可验证交付物为入口。",
        },
        {
            "risks": [
                {"risk": "输出无效", "impact": "高", "action": "约束生成并执行确定性验证", "owner": "算法组"},
                {"risk": "服务延迟", "impact": "中", "action": "限制解码并启用缓存", "owner": "平台组"},
                {"risk": "目标漂移", "impact": "中", "action": "持续监控并保留回滚", "owner": "产品组"},
            ],
            "rule": "优先处理高影响风险，每项风险必须绑定负责人和可验证措施。",
        },
    ]
    prs, _warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=["cover", "ablation_matrix", "workpackage_matrix", "risk_action_table"],
        content_payloads=payloads,
        image_paths={},
    )
    ablation = prs.slides[1]
    assert not named(ablation, "ablation-guide")
    assert named(ablation, "ablation-callout")[0].height / 914400 >= 4.9
    assert named(ablation, "ablation-grid-r2c1")[0].line.fill.type == MSO_FILL_TYPE.BACKGROUND
    assert min_font(named(ablation, "ablation-callout")[0]) >= 14

    workpackages = prs.slides[2]
    assert named(workpackages, "wp-grid-r2c3")[0].width > named(workpackages, "wp-grid-r2c1")[0].width
    assert named(workpackages, "wp-grid-r2c1")[0].line.fill.type == MSO_FILL_TYPE.BACKGROUND
    assert named(workpackages, "wp-note")[0].height / 914400 >= 0.6

    risk = prs.slides[3]
    assert named(risk, "risk-grid-r1c1")[0].text == "风险"
    assert named(risk, "risk-grid-r2c1")[0].line.fill.type == MSO_FILL_TYPE.BACKGROUND
    assert named(risk, "risk-rule")[0].height / 914400 >= 0.6


def variety_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    specs = [spec("Cover"), spec("Pipeline one"), spec("Challenge one"), spec("Pipeline two"), spec("Challenge two")]
    pipeline = {
        "input": "Context",
        "steps": [{"label": "Plan", "detail": "Select a bounded use case"}, {"label": "Build", "detail": "Build a prototype"}],
        "output": "Validated output",
    }
    challenge = {
        "risks": [
            {"risk": "Invalid output", "impact": "Low trust", "mitigation": "Validate against the catalog"},
            {"risk": "Latency", "impact": "Slow serving", "mitigation": "Use bounded decoding"},
        ],
    }
    prs, warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=["cover", "method_pipeline", "challenge_map", "method_pipeline", "challenge_map"],
        content_payloads=[{}, pipeline, challenge, pipeline, challenge],
        image_paths={},
    )
    assert named(prs.slides[3], "tpl-title-17"), "repeated pipeline should use technical_route"
    assert visible_text(prs.slides[3]).count("Plan") == 1
    assert named(prs.slides[4], "tpl-title-20"), "repeated challenge should use risk_action_table"
    assert not named(prs.slides[4], "risk-grid-r1c4"), "three-field risk table should remove empty columns"
    assert sum("替补版式" in warning for warning in warnings) >= 2


def deck_rhythm_and_quality_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    template_id = "scut" if "华工" in template_name else "hitsz"
    layouts = [
        "cover",
        "basic_content",
        "basic_content",
        "summary_takeaways",
        "method_pipeline",
        "method_pipeline",
        "benchmark_metrics",
        "benchmark_metrics",
    ]
    payloads = [
        {},
        {"points": ["Frame the decision", "State the evidence", "Name the next action"]},
        {"points": ["Clarify the boundary", "Compare alternatives", "Record the decision"]},
        {"cards": [{"title": "Finding", "points": ["Evidence supports the change"]}, {"title": "Constraint", "points": ["The pilot is still limited"]}]},
        {"input": "Request", "steps": ["Interpret", "Ground", "Validate"], "output": "Decision"},
        {"input": "Evidence", "steps": ["Prepare", "Model", "Test"], "output": "Result"},
        {"kpis": [{"value": "91%", "label": "Quality"}, {"value": "170ms", "label": "Latency"}]},
        {"kpis": [{"value": "94%", "label": "Validity"}, {"value": "0.8x", "label": "Cost"}]},
    ]
    specs = [spec(payload.get("title") or f"Slide {index}") for index, payload in enumerate(payloads, start=1)]
    prs, warnings = render_slide_library_presentation(
        template_path=template,
        template_id=template_id,
        specs=specs,
        layout_ids=layouts,
        content_payloads=payloads,
        image_paths={},
    )
    expected_markers = [
        "tpl-title-2",
        "tpl-title-14",
        "tpl-title-12",
        "tpl-title-5",
        "tpl-title-17",
        "tpl-title-7",
        "tpl-title-8",
    ]
    for slide, marker in zip(list(prs.slides)[1:], expected_markers):
        assert named(slide, marker), f"deck rhythm should select {marker}"
    assert sum("页面节奏" in warning and "替补版式" in warning for warning in warnings) >= 4

    dense_payloads = [
        {},
        {"rows": [["Risk A", "High", "Mitigate", "Owner"]]},
        {"rows": [["WP1", "Owner", "Task", "Output", "Metric", "Q1"]]},
    ]
    dense_specs = [
        SlideSpec(title="Cover", content_spec={}),
        SlideSpec(title="Risks", content_spec=dense_payloads[1]),
        SlideSpec(title="Work packages", content_spec=dense_payloads[2]),
    ]
    dense_prs, dense_warnings = render_slide_library_presentation(
        template_path=template,
        template_id=template_id,
        specs=dense_specs,
        layout_ids=["cover", "risk_action_table", "workpackage_matrix"],
        content_payloads=dense_payloads,
        image_paths={},
    )
    assert any("连续使用高密度页面" in warning for warning in dense_warnings)
    dense_out = TMP / f"dense-rhythm-{template.stem}.pptx"
    dense_prs.save(dense_out)
    dense_issues = _ppt_text_quality_issues(dense_out, dense_specs)
    assert any("连续为高密度页面" in issue for issue in dense_issues)
    dense_out.unlink(missing_ok=True)

    modified_template = TMP / f"safe-area-{template.stem}.pptx"
    source = Presentation(template)
    summary_slide = next(
        slide
        for slide in source.slides
        if any(shape.name.startswith("tpl-title-") and shape.text.strip() == "summary_takeaways" for shape in slide.shapes)
    )
    sentinel = summary_slide.shapes.add_textbox(Inches(-0.45), Inches(0.25), Inches(1.2), Inches(0.4))
    sentinel.name = "summary-quality-sentinel"
    sentinel.text = "Quality sentinel"
    source.save(modified_template)
    repaired, repair_warnings = render_slide_library_presentation(
        template_path=modified_template,
        template_id=template_id,
        specs=[spec("Cover"), spec("Summary")],
        layout_ids=["cover", "summary_takeaways"],
        content_payloads=[{}, {"cards": ["Finding"], "next_step": "Act"}],
        image_paths={},
    )
    repaired_sentinel = named(repaired.slides[1], "summary-quality-sentinel")[0]
    assert repaired_sentinel.left / 914400 >= 0.35
    assert repaired_sentinel.top / 914400 >= 1.05
    assert any("超出正文安全区" in warning for warning in repair_warnings)
    modified_template.unlink(missing_ok=True)


def adaptive_density_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    dense_image = TMP / "ppt-dense-regression-image.png"
    dense_image.parent.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (1280, 720), "white")
    draw = ImageDraw.Draw(canvas)
    for row in range(8):
        for col in range(12):
            x, y = 25 + col * 103, 25 + row * 82
            draw.rectangle((x, y, x + 78, y + 52), outline=(30 + row * 15, 80 + col * 8, 145), width=3)
            draw.line((x + 8, y + 18, x + 68, y + 18), fill=(40, 80, 120), width=2)
            draw.line((x + 8, y + 32, x + 58, y + 32), fill=(70, 110, 145), width=2)
    canvas.save(dense_image)
    specs = [spec("Cover"), spec("Dense visual"), spec("Text only"), spec("Summary"), spec("Ablation"), spec("Benchmark")]
    payloads = [
        {},
        {"points": ["First supporting point", "Second supporting point", "Third supporting point"]},
        {"points": ["Question one needs context", "Question two needs evidence", "Question three needs evaluation"]},
        {"cards": [{"title": "Scope", "points": ["Generate grounded outputs"]}, {"title": "Control", "points": ["Validate constraints"]}]},
        {
            "headers": ["Component", "Contribution", "Failure", "Test"],
            "rows": [["Retriever", "Evidence", "Invalid output", "Remove retrieval"]],
            "conclusion": "Test each component",
        },
        {
            "kpis": [
                {"value": "HR / NDCG", "label": "Relevance", "note": "Measure ranking utility on held-out interactions"},
                {"value": "Validity / Coverage", "label": "Generation quality", "note": "Check whether outputs are valid and sufficiently grounded"},
                {"value": "Diversity / Novelty", "label": "Discovery", "note": "Quantify catalog breadth and non-obvious recommendations"},
            ],
            "protocol": "Use fixed splits and identical serving budgets",
        },
    ]
    prs, _ = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=["cover", "basic_content", "basic_content", "summary_takeaways", "ablation_matrix", "benchmark_metrics"],
        content_payloads=payloads,
        image_paths={2: dense_image},
    )
    dense_slide = prs.slides[1]
    dense_picture = named(dense_slide, "janus-basic-image-image")[0]
    assert dense_picture.width / 914400 >= 9.0
    assert not named(dense_slide, "basic-text-card")
    notes_text = dense_slide.notes_slide.notes_text_frame.text
    assert "First supporting point" in notes_text and "Third supporting point" in notes_text
    text_only = prs.slides[2]
    assert named(text_only, "tpl-title-14"), "short repeated text should become a summary page to improve deck rhythm"
    assert min_font(named(text_only, "summary-card1")[0]) >= 15
    summary_text = visible_text(prs.slides[3])
    assert "Scope" in summary_text and "title：" not in summary_text and "points：" not in summary_text
    assert not named(prs.slides[4], "ablation-grid-r1c5")
    assert not named(prs.slides[4], "ablation-guide")
    assert not named(prs.slides[5], "bench-table-")
    assert named(prs.slides[5], "bench-kpi1")[0].height / 914400 >= 3.5
    dense_image.unlink(missing_ok=True)


def deterministic_text_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    specs = [
        spec("Cover"),
        spec("Project target"),
        spec("Data and scenario construction"),
        spec("Evaluation"),
        spec("Risks"),
        spec("Roadmap"),
    ]
    target_payload = {
        "title": specs[1].title,
        "objective": "Represent intent while preserving production control",
        "nodes": [
            {"label": "Candidate dependence", "detail": "A relevant item cannot be selected if it never enters the candidate set."},
            {"label": "Fragmented objectives", "detail": "Retrieval and ranking may be optimized separately."},
            {"label": "Sparse generalization", "detail": "Cold-start and long-tail behavior remain difficult."},
            {"label": "Complex output needs", "detail": "Bundles and constrained choices require structured outputs."},
        ],
        "row": ["Does it improve relevance?", "Does it generalize?", "Are outputs valid?", "Can it be served efficiently?"],
    }
    domain_payload = {
        "title": specs[2].title,
        "layers": ["Raw observations", "Training representation", "Learning objective", "Evaluation scenarios", "Feedback"],
        "nodes": [
            {"label": "Interaction events", "detail": "Views, clicks, purchases, skips, dwell time, and session boundaries"},
            {"label": "Catalog encoding", "detail": "Stable identifiers, semantic codes, text, images, and structured attributes"},
            {"label": "Training examples", "detail": "Next-token prediction, sequence completion, denoising, or preference alignment"},
            {"label": "Scenario slices", "detail": "Head, long-tail, cold-start, sparse-history, and multi-intent users"},
            {"label": "Outcome signals", "detail": "User response, constraint violations, exposure patterns, and drift"},
        ],
        "legend": ["Historical data may contain exposure and selection bias", "Scenario-specific evidence is required"],
    }
    evaluation_payload = {
        "title": specs[3].title,
        "scope": "Compare against strong baselines across representative scenario slices.",
        "kpis": [
            {"value": "Recall@K / NDCG@K", "label": "Offline relevance", "note": "Report by head, long-tail, cold-start, and sparse-history slices."},
            {"value": "CTR / CVR / retention", "label": "Online value", "note": "Interpret together with satisfaction, diversity, and exposure guardrails."},
            {"value": "Latency / cost", "label": "System health", "note": "Include catalog validity, policy compliance, and fallback frequency."},
        ],
        "status_rows": [
            {"item": "Baselines", "status": "Required", "evidence": "Strong retrieval-ranking and sequential models"},
            {"item": "Scenario slicing", "status": "Required", "evidence": "Head, tail, cold-start, sparse, and multi-intent cases"},
            {"item": "Ablations", "status": "Required", "evidence": "Representation, generation constraints, and hybrid reranking"},
        ],
    }
    risk_payload = {
        "title": specs[4].title,
        "risks": [
            {"risk": "Invalid output", "impact": "Broken recommendations", "action": "Use constrained decoding and deterministic validation", "owner": "TBD"},
            {"risk": "Popularity bias", "impact": "Weak long-tail coverage", "action": "Apply slice metrics and exposure controls", "owner": "TBD"},
            {"risk": "Unsafe content", "impact": "Compliance exposure", "action": "Ground outputs in verified facts", "owner": "TBD"},
            {"risk": "Serving cost", "impact": "Unsustainable economics", "action": "Use bounded decoding and caching", "owner": "TBD"},
            {"risk": "Objective drift", "impact": "Reduced user value", "action": "Monitor drift and preserve rollback", "owner": "TBD"},
        ],
    }
    roadmap_payload = {
        "title": specs[5].title,
        "milestones": [
            {"time": f"Phase {index}", "title": title, "deliverable": f"Detailed deliverable {index}"}
            for index, title in enumerate(["Define scope", "Build prototype", "Offline evidence", "Online pilot", "Scale safely"], start=1)
        ],
    }
    prs, _ = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=["cover", "project_target_map", "domain_object_map", "evaluation_dashboard", "risk_action_table", "milestone_roadmap"],
        content_payloads=[{}, target_payload, domain_payload, evaluation_payload, risk_payload, roadmap_payload],
        image_paths={},
    )
    target = prs.slides[1]
    assert all("label" not in shape.text.lower() and "detail" not in shape.text.lower() for shape in named(target, "target-node"))
    assert min_font(named(target, "target-grid-r1c1")[0]) >= 12
    assert "candidate set" in target.notes_slide.notes_text_frame.text
    slide = prs.slides[2]
    nodes = named(slide, "domain-node")
    assert len(nodes) == 5
    assert all(shape.width / 914400 >= 2.0 for shape in nodes)
    assert all(min_font(shape) >= 14 for shape in nodes)
    assert all("label" not in shape.text.lower() and "detail" not in shape.text.lower() for shape in nodes)
    assert [shape.text for shape in nodes] == [item["label"] for item in domain_payload["nodes"]]
    assert "session boundaries" in slide.notes_slide.notes_text_frame.text
    evaluation = prs.slides[3]
    assert not named(evaluation, "eval-bar")
    assert min_font(named(evaluation, "eval-status-grid-r2c3")[0]) >= 12
    risk = prs.slides[4]
    assert not named(risk, "risk-grid-r1c5")
    assert named(risk, "risk-grid-r6c1")[0].top + named(risk, "risk-grid-r6c1")[0].height <= named(risk, "risk-rule")[0].top
    roadmap = prs.slides[5]
    assert not named(roadmap, "roadmap-card-5")
    assert not named(roadmap, "roadmap-dot-5")
    assert "Detailed deliverable 5" in roadmap.notes_slide.notes_text_frame.text
    for rendered_slide in prs.slides:
        for shape in rendered_slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                assert shape.text_frame.auto_size != MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def schema_adaptation_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    payloads = [
        {},
        {
            "title": "Why this project matters",
            "objective": "Create measurable value",
            "nodes": [
                {"label": "Need", "detail": "Unmet demand"},
                {"label": "Capability", "detail": "Technical readiness"},
                {"label": "Evidence", "detail": "Verifiable outcomes"},
                {"label": "Control", "detail": "Operational guardrails"},
            ],
        },
        {
            "title": "Object topology",
            "layers": [
                {"label": "Context", "detail": "User and session information"},
                {"label": "Control", "detail": "Policy and validation"},
            ],
            "nodes": [
                {"label": "Input", "detail": "Observed request"},
                {"label": "Model", "detail": "Structured decision"},
                {"label": "Output", "detail": "Validated result"},
            ],
        },
        {
            "title": "Evaluation system",
            "kpis": [
                {"value": "Relevance", "label": "Preference fit", "note": "Does the output match the current need?"},
                {"value": "Utility", "label": "Task completion", "note": "Does it produce a useful outcome?"},
                {"value": "Trust", "label": "Grounding and safety", "note": "Is it valid and supported?"},
            ],
            "status_rows": [
                {"dimension": "Offline evidence", "signals": "Quality, diversity, and constraint satisfaction", "purpose": "Model screening"},
                {"dimension": "Online evidence", "signals": "Task success and correction", "purpose": "Product validation"},
                {"dimension": "System evidence", "signals": "Latency, cost, and recovery", "purpose": "Deployment readiness"},
            ],
        },
        {
            "title": "Risk controls",
            "rows": [
                {"risk": "Invalid output", "trigger": "Unconstrained generation", "impact": "Broken experience", "action": "Validate against the catalog"},
                {"risk": "Latency", "trigger": "Long decoding", "impact": "Slow service", "action": "Use bounded generation"},
            ],
        },
        {
            "title": "Three product patterns",
            "layers": ["Low autonomy", "High autonomy"],
            "nodes": [
                {"label": "Assistant", "detail": "Explain and compare"},
                {"label": "Co-creation", "detail": "Refine with the user"},
                {"label": "Agent", "detail": "Execute approved actions"},
            ],
        },
        {
            "title": "Evidence gates",
            "kpis": [
                {"value": "01", "label": "User value", "note": "Faster discovery and clearer decisions"},
                {"value": "02", "label": "Business value", "note": "Differentiation and efficient service"},
                {"value": "03", "label": "System value", "note": "Reusable and controlled capability"},
            ],
        },
    ]
    layout_ids = ["cover", "project_target_map", "domain_object_map", "evaluation_dashboard", "risk_action_table", "domain_object_map", "result_big_numbers"]
    specs = [
        SlideSpec(title=(payload.get("title") or "Cover"), visual=f"layout_id: {layout}", content_spec=payload)
        for payload, layout in zip(payloads, layout_ids)
    ]
    prs, warnings = render_slide_library_presentation(
        template_path=template,
        template_id="scut" if "华工" in template_name else "hitsz",
        specs=specs,
        layout_ids=layout_ids,
        content_payloads=payloads,
        image_paths={},
    )
    assert not named(prs.slides[1], "target-grid-")
    assert visible_text(prs.slides[2]).count("label：") == 0 and visible_text(prs.slides[2]).count("detail：") == 0
    assert len(named(prs.slides[2], "domain-edge")) == 2
    assert "Required signals" in named(prs.slides[3], "eval-status-grid-r1c2")[0].text
    assert min_font(named(prs.slides[3], "eval-kpi1-label")[0]) >= 12
    assert not named(prs.slides[4], "risk-grid-r1c5")
    assert named(prs.slides[4], "risk-grid-r1c2")[0].text == "Trigger"
    assert named(prs.slides[4], "risk-grid-r2c4")[0].text == "Validate against the catalog"
    assert named(prs.slides[5], "tpl-title-12"), "pattern/category content should adapt to evidence_grid"
    assert any("分类/对比" in warning for warning in warnings)
    assert min_font(named(prs.slides[6], "big-kpi2-label")[0]) >= 12
    assert named(prs.slides[6], "big-kpi1")[0].width / 914400 >= 5.0
    assert abs(named(prs.slides[6], "big-kpi2")[0].left - named(prs.slides[6], "big-kpi3")[0].left) <= 2
    out = TMP / f"schema-adaptation-{template.stem}.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    issues = _ppt_text_quality_issues(out, specs)
    blocking_issues = [issue for issue in issues if "连续为高密度页面" not in issue]
    assert not blocking_issues, blocking_issues
    assert any("连续为高密度页面" in issue for issue in issues)
    out.unlink(missing_ok=True)


def architecture_layout_checks(template_name: str) -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / template_name
    template_id = "scut" if "华工" in template_name else "hitsz"

    def render_one(layout_id: str, payload: dict) -> object:
        specs = [spec("Architecture cover"), spec(payload.get("title") or layout_id)]
        prs, _ = render_slide_library_presentation(
            template_path=template,
            template_id=template_id,
            specs=specs,
            layout_ids=["cover", layout_id],
            content_payloads=[{}, payload],
            image_paths={},
        )
        return prs.slides[1]

    loop_three = render_one(
        "method_loop",
        {
            "title": "Three-stage learning loop",
            "objective": "Improve verified user value",
            "steps": ["Plan", "Execute", "Learn"],
            "guardrail": "Stop or revise when evidence falls below the threshold",
        },
    )
    loop_nodes = named(loop_three, "loop-node-")
    loop_arrows = named(loop_three, "loop-a")
    assert len(loop_nodes) == 3 and len(loop_arrows) == 3
    assert loop_nodes[0].top < loop_nodes[1].top and abs(loop_nodes[1].top - loop_nodes[2].top) <= 2
    for index, arrow in enumerate(loop_arrows):
        start, end = connector_endpoints(arrow)
        source = loop_nodes[index]
        target = loop_nodes[(index + 1) % len(loop_nodes)]
        assert point_on_shape_border(start, source)
        assert point_on_shape_border(end, target)
    assert named(loop_three, "loop-center")[0].width / 914400 >= 3.5
    assert named(loop_three, "loop-guardrail")[0].width / 914400 >= 11.5

    loop_four = render_one(
        "method_loop",
        {
            "title": "Four-stage control loop",
            "objective": "Controlled iteration",
            "steps": ["Plan", "Act", "Reflect", "Observe"],
        },
    )
    assert len(named(loop_four, "loop-node-")) == 4
    assert len(named(loop_four, "loop-a")) == 4

    route = render_one(
        "technical_route",
        {
            "title": "Adaptive technical route",
            "stages": [
                {"title": "Foundation", "task": "Prepare data and constraints"},
                {"title": "Core method", "task": "Build and validate the central method"},
                {"title": "Demonstration", "task": "Integrate, evaluate, and demonstrate"},
            ],
            "dependency": "Shared data contracts and evaluation gates connect every stage",
        },
    )
    route_headers = named(route, "route-stage-")
    assert len(route_headers) == 3
    assert len(named(route, "route-task-")) == 3
    assert len(named(route, "route-arrow-")) == 2
    assert all(shape.height / 914400 >= 2.9 for shape in named(route, "route-task-"))
    for index, arrow in enumerate(named(route, "route-arrow-")):
        start, end = connector_endpoints(arrow)
        assert point_on_shape_border(start, route_headers[index])
        assert point_on_shape_border(end, route_headers[index + 1])
    assert named(route, "route-dependency")[0].top / 914400 >= 5.9

    target = render_one(
        "project_target_map",
        {
            "title": "Traceable project targets",
            "objective": "Deliver measurable, reusable outcomes",
            "nodes": ["Scientific question", "Technical system", "Demonstration evidence"],
        },
    )
    target_root = named(target, "target-root")[0]
    target_nodes = named(target, "target-node")
    assert target_root.width / 914400 >= 11.5 and len(target_nodes) == 3
    assert not named(target, "target-grid-")
    for link, node in zip(named(target, "target-link"), target_nodes):
        root_endpoint, node_endpoint = connector_endpoints(link)
        assert point_on_shape_border(root_endpoint, target_root)
        assert point_on_shape_border(node_endpoint, node)

    domain = render_one(
        "domain_object_map",
        {
            "title": "Domain object topology",
            "layers": ["Context layer", "Decision layer", "Evidence layer"],
            "nodes": ["Context", "Method", "Control", "Evidence"],
            "legend": "Lines express the minimum dependency structure",
        },
    )
    domain_nodes = named(domain, "domain-node")
    domain_edges = named(domain, "domain-edge")
    domain_pairs = ((0, 1), (0, 2), (1, 3), (2, 3))
    assert len(domain_nodes) == 4 and len(domain_edges) == 4
    for edge, (source_index, target_index) in zip(domain_edges, domain_pairs):
        start, end = connector_endpoints(edge)
        assert point_on_shape_border(start, domain_nodes[source_index])
        assert point_on_shape_border(end, domain_nodes[target_index])
    assert named(domain, "domain-legend")[0].width / 914400 >= 8.5

    roadmap = render_one(
        "milestone_roadmap",
        {
            "title": "Milestone status roadmap",
            "milestones": [
                {"time": "M1", "title": "Scope", "status": "completed"},
                {"time": "M2", "title": "Prototype", "state": "current"},
                {"time": "M3", "title": "Pilot", "status": "pending"},
                {"time": "M4", "title": "Scale", "status": "at risk"},
            ],
            "note": "Each milestone closes with evidence and an explicit next gate",
        },
    )
    dots = named(roadmap, "roadmap-dot-")
    assert len(named(roadmap, "roadmap-card-")) == 4
    assert len(dots) == 4 and len(named(roadmap, "roadmap-link-")) == 4
    centers = [float(dot.left + dot.width / 2) for dot in dots]
    gaps = [centers[index + 1] - centers[index] for index in range(len(centers) - 1)]
    assert max(gaps) - min(gaps) <= 3
    assert [str(dot.fill.fore_color.rgb) for dot in dots] == ["34855B", "2F6FAE", "9EAAB4", "E08737"]
    assert named(roadmap, "roadmap-note")[0].width / 914400 >= 11.5


def neutral_template_checks() -> None:
    template = ROOT / "assets" / "departments" / "ppt_department" / "templates" / "通用多功能PPT模板.pptx"
    assert not validate_slide_library(template)
    decision = resolve_ppt_template(ASSET_ROOT, "none")
    assert decision.path is not None and decision.path.name == template.name
    specs = [spec("Neutral cover", "No school branding"), spec("Why now"), spec("Constraints")]
    payloads = [
        {},
        {
            "current": {"label": "Today", "points": ["Fragmented workflow"]},
            "target": {"label": "Future", "points": ["Reusable workflow"]},
            "gap": "A shared operating model",
        },
        {
            "nodes": [
                {"label": "Value", "detail": "Useful outcome"},
                {"label": "Evidence", "detail": "Grounded claim"},
                {"label": "Control", "detail": "Bounded action"},
                {"label": "Operations", "detail": "Reliable service"},
            ],
            "objective": "Trustworthy delivery",
        },
    ]
    prs, warnings = render_slide_library_presentation(
        template_path=template,
        template_id="none",
        specs=specs,
        layout_ids=["cover", "motivation_compare", "challenge_map"],
        content_payloads=payloads,
        image_paths={},
    )
    assert not warnings
    assert named(prs.slides[0], "neutral-cover-kicker")
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for slide in prs.slides for shape in slide.shapes)
    assert not re.search(r"华南理工|哈尔滨工业|SCUT|HITSZ", " ".join(visible_text(slide) for slide in prs.slides), re.IGNORECASE)
    assert not named(prs.slides[2], "challenge-note")
    assert named(prs.slides[1], "tpl-title-")[0].line.fill.type == MSO_FILL_TYPE.BACKGROUND


for filename in ("华工多功能PPT模板.pptx", "哈工深多功能PPT模板.pptx"):
    run(filename)
    font_and_image_checks(filename)
    image_caption_and_density_checks(filename)
    data_visual_hierarchy_checks(filename)
    semantic_emphasis_checks(filename)
    matrix_visual_hierarchy_checks(filename)
    variety_checks(filename)
    deck_rhythm_and_quality_checks(filename)
    adaptive_density_checks(filename)
    deterministic_text_checks(filename)
    schema_adaptation_checks(filename)
    architecture_layout_checks(filename)

neutral_template_checks()


def semantic_contract_repair_checks() -> None:
    specs = [
        SlideSpec(title="Cover", message="Multimodal recommendation", visual="layout_id: basic_content"),
        SlideSpec(
            title="从 ID 推荐到多模态语义推荐",
            message="现有方法依赖 ID • 目标是理解图像与文本内容",
            visual="layout_id: motivation_compare",
            content_spec={"current": "依赖 ID", "target": "依赖 ID"},
        ),
        SlideSpec(
            title="编码—融合—对齐—排序",
            message="四阶段技术流程",
            visual="layout_id: method_pipeline",
            content_spec={"steps": ["编码—融合—对齐—排序"]},
        ),
        SlideSpec(
            title="训练、服务、反馈闭环",
            message="训练后上线，并根据反馈更新",
            visual="layout_id: method_loop",
            content_spec={"steps": ["训练", "训练", "服务、反馈"]},
        ),
        SlideSpec(
            title="代表模型比较",
            message="VBPR 引入视觉特征 • 图模型传播多模态信号",
            visual="layout_id: leaderboard_table",
            content_spec={"rows": [["1", "VBPR"]]},
        ),
        SlideSpec(
            title="评测设计",
            message="同时关注排序质量与长尾覆盖",
            visual="layout_id: benchmark_metrics",
            content_spec={"kpis": [{"value": "Recall@K", "label": "排序质量"}]},
        ),
        SlideSpec(
            title="消融分析",
            message="分别验证视觉、文本与融合模块",
            visual="layout_id: ablation_matrix",
            content_spec={"rows": [["完整方案", "✓"]]},
        ),
        SlideSpec(
            title="总结",
            message="内容理解改善冷启动 • 评测需要兼顾长尾",
            visual="layout_id: summary_takeaways",
            content_spec={},
        ),
    ]
    repaired = ppt_service._repair_semantic_content_specs("帮我生成一个介绍多模态推荐的 ppt", specs)
    motivation = repaired[1].content_spec
    assert ppt_service._semantic_text(motivation["current"]) != ppt_service._semantic_text(motivation["target"])
    assert len(repaired[2].content_spec["steps"]) == 4
    assert repaired[2].content_spec["steps"] == ["编码", "融合", "对齐", "排序"]
    assert 3 <= len(repaired[3].content_spec["steps"]) <= 4
    assert len({ppt_service._semantic_text(item) for item in repaired[3].content_spec["steps"]}) == len(repaired[3].content_spec["steps"])
    assert ppt_service._spec_layout_id(repaired[4]) in {"basic_content", "evidence_grid"}
    assert ppt_service._spec_layout_id(repaired[5]) in {"basic_content", "evidence_grid"}
    assert ppt_service._spec_layout_id(repaired[6]) in {"basic_content", "evidence_grid"}
    summary_items = repaired[7].content_spec["cards"]
    assert len(summary_items) == 3
    assert len({ppt_service._semantic_text(item) for item in summary_items}) == 3


semantic_contract_repair_checks()


def image_prompt_diversity_checks() -> None:
    academic = SimpleNamespace(
        style_id="academic_report",
        label="学术汇报风",
        prompt="publication-grade scientific editorial imagery",
        palette={"ink": "0F172A", "accent": "0F766E", "accent2": "2563EB", "bg": "F8FAFC"},
    )
    project = SimpleNamespace(
        style_id="major_project",
        label="重大项目风",
        prompt="cinematic engineering and industrial editorial imagery",
        palette={"ink": "111827", "accent": "1D4ED8", "accent2": "D97706", "bg": "F7FAFC"},
    )
    custom = SimpleNamespace(
        style_id="custom_ink_wash",
        label="水墨风",
        prompt="Chinese ink-wash painting with expressive negative space",
        palette={"ink": "111827", "accent": "2563EB", "accent2": "0F766E", "bg": "F8FAFC"},
        created=True,
    )
    medical = SlideSpec(
        title="蛋白质药物递送机制",
        message="纳米载体穿过细胞膜并在靶组织释放药物",
        visual="layout_id: media_showcase\n生成插图：细胞膜附近的纳米载体递送场景",
    )
    industrial = SlideSpec(
        title="智能制造平台落地",
        message="机器人、产线设备和质量检测系统形成闭环",
        visual="layout_id: media_showcase\n生成插图：真实工厂中的机器人协同生产场景",
    )
    academic_prompt = ppt_service._slide_image_prompt(medical, academic)
    project_prompt = ppt_service._slide_image_prompt(industrial, project)
    custom_prompt = ppt_service._slide_image_prompt(industrial, custom)
    assert "Life-science world" in academic_prompt
    assert "Publication-grade scientific editorial" in academic_prompt
    assert "Engineering world" in project_prompt
    assert "Cinematic engineering and industrial editorial" in project_prompt
    assert "3-5 distinct zones/nodes/callouts" not in academic_prompt
    assert "Suggested short labels/themes" not in academic_prompt
    assert "render no words, letters, numbers" in academic_prompt
    assert "do not inherit the default corporate blue palette" in custom_prompt
    assert "Chinese ink-wash painting" in custom_prompt
    assert academic_prompt != project_prompt


image_prompt_diversity_checks()


def imagegen_cache_and_parallel_checks() -> None:
    cache_dir = TMP / "imagegen-cache"
    sessions = ["regression-imagegen-first", "regression-imagegen-cached"]
    for session in sessions:
        output = ROOT / "outputs" / "ppt_department" / session
        if output.exists():
            import shutil
            shutil.rmtree(output)

    specs = [
        SlideSpec(title="Cover", message="", visual="layout_id: cover"),
        SlideSpec(title="Generated scene A", message="", visual="layout_id: media_showcase\n生成插图"),
        SlideSpec(title="Generated scene B", message="", visual="layout_id: basic_content\n生成插图"),
        SlideSpec(title="Summary", message="", visual="layout_id: summary_takeaways"),
    ]
    style = SimpleNamespace(prompt="clean academic supporting visual")
    active = 0
    peak_active = 0
    calls: list[str] = []
    sizes: list[str] = []
    lock = threading.Lock()

    def fake_generate_image_direct(*, prompt: str, out_path: Path, **_kwargs) -> None:
        nonlocal active, peak_active
        with lock:
            active += 1
            peak_active = max(peak_active, active)
            calls.append(prompt)
            sizes.append(str(_kwargs.get("size") or ""))
        try:
            time.sleep(0.08)
            Image.new("RGB", (320, 180), (36, 88, 142)).save(out_path)
        finally:
            with lock:
                active -= 1

    env = {
        "OPENAI_API_KEY": "",
        "OPENAI_IMAGE_API_KEY": "ppt-regression-placeholder",
        "JANUS_PPT_IMAGEGEN_MAX": "2",
        "JANUS_PPT_IMAGEGEN_ATTEMPTS": "1",
        "JANUS_PPT_IMAGEGEN_CONCURRENCY": "2",
        "JANUS_PPT_IMAGEGEN_CACHE": "1",
        "JANUS_PPT_IMAGEGEN_CACHE_DIR": str(cache_dir),
    }
    first_events: list[dict[str, object]] = []
    with patch.dict(os.environ, env, clear=False), patch.object(
        ppt_service._image_generation_module,
        "_ppt_imagegen_env",
        return_value={"OPENAI_IMAGE_API_KEY": "ppt-regression-placeholder"},
    ), patch.object(
        ppt_service,
        "_ppt_generate_image_direct",
        side_effect=fake_generate_image_direct,
    ):
        first_images, first_errors = ppt_service._generate_slide_images(
            root=ROOT,
            session_id=sessions[0],
            specs=specs,
            style=style,
            enable_imagegen=True,
            selected_template="none",
            on_progress=first_events.append,
        )
    assert not first_errors
    assert sorted(first_images) == [2, 3]
    assert len(calls) == 2
    assert sizes == ["1536x1024", "1536x1024"]
    assert peak_active == 2, "independent slide images should use the configured parallelism"
    completed = [event for event in first_events if event.get("status") == "completed"]
    assert len(completed) == 2
    assert max(int(event.get("phase_current") or 0) for event in completed) == 2
    assert all(int(event.get("concurrency") or 0) == 2 for event in completed)

    calls.clear()
    cached_events: list[dict[str, object]] = []
    with patch.dict(os.environ, env, clear=False), patch.object(
        ppt_service._image_generation_module,
        "_ppt_imagegen_env",
        return_value={"OPENAI_IMAGE_API_KEY": "ppt-regression-placeholder"},
    ), patch.object(
        ppt_service,
        "_ppt_generate_image_direct",
        side_effect=fake_generate_image_direct,
    ):
        cached_images, cached_errors = ppt_service._generate_slide_images(
            root=ROOT,
            session_id=sessions[1],
            specs=specs,
            style=style,
            enable_imagegen=True,
            selected_template="none",
            on_progress=cached_events.append,
        )
    assert not cached_errors
    assert sorted(cached_images) == [2, 3]
    assert not calls, "the second identical request must reuse generated-image cache entries"
    cached = [event for event in cached_events if event.get("status") == "cached"]
    assert len(cached) == 2
    assert int(cached[-1].get("cache_hits") or 0) == 2

    import shutil
    shutil.rmtree(cache_dir, ignore_errors=True)
    for session in sessions:
        shutil.rmtree(ROOT / "outputs" / "ppt_department" / session, ignore_errors=True)


imagegen_cache_and_parallel_checks()


def image_api_key_fallback_check() -> None:
    output = TMP / "image-api-key-fallback.png"
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (22, 88, 144)).save(buffer, format="PNG")
    body = json.dumps({"data": [{"b64_json": base64.b64encode(buffer.getvalue()).decode("ascii")}]}).encode("utf-8")
    authorization: list[str] = []

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self):
            return body

    def fake_urlopen(request, timeout=0):
        authorization.append(request.headers.get("Authorization", ""))
        assert timeout > 0
        return FakeResponse()

    with patch.object(ppt_service._image_generation_module.urllib.request, "urlopen", side_effect=fake_urlopen):
        ppt_service._ppt_generate_image_direct(
            root=ROOT,
            prompt="A specific editorial illustration without text",
            out_path=output,
            quality="medium",
            size="1536x1024",
            env={"OPENAI_IMAGE_API_KEY": "", "OPENAI_API_KEY": "fallback-key"},
        )
    assert authorization == ["Bearer fallback-key"]
    assert output.is_file()
    output.unlink(missing_ok=True)


image_api_key_fallback_check()

selection_specs = [
    SlideSpec(
        title=f"{'Application context' if layout in {'basic_content', 'case_gallery'} else 'Slide'} {index}",
        message="",
        visual=f"layout_id: {layout}",
    )
    for index, layout in enumerate([
        "basic_content", "motivation_compare", "evidence_grid", "basic_content", "method_pipeline", "method_loop",
        "benchmark_metrics", "leaderboard_table", "case_gallery", "challenge_map", "evidence_grid", "summary_takeaways",
    ], start=1)
]
assert _selected_image_slide_indices(selection_specs, 12, include_cover=False)[:2] == [4, 9]

spread_specs = [
    SlideSpec(
        title=f"{'Problem context' if layout == 'basic_content' else 'Application evidence' if layout == 'evidence_grid' else 'Slide'} {index}",
        message="",
        visual=f"layout_id: {layout}",
    )
    for index, layout in enumerate([
        "basic_content", "motivation_compare", "basic_content", "evidence_grid", "method_pipeline",
        "method_loop", "challenge_map", "evidence_grid", "benchmark_metrics", "summary_takeaways",
    ], start=1)
]
assert _selected_image_slide_indices(spread_specs, 2, include_cover=False) == [3, 8]

source_paths: list[Path] = []
for index in range(1, 4):
    path = TMP / f"paper-figure-{index}.png"
    Image.new("RGB", (320, 180), (30 * index, 70, 140)).save(path)
    source_paths.append(path)
source_specs = [
    SlideSpec(title="Cover", message=""),
    SlideSpec(
        title="Source evidence",
        message="",
        visual="layout_id: evidence_grid\n使用附件原图：图1、图2、图3",
    ),
]
source_map = _map_source_visuals_to_slides(source_specs, source_paths)
assert source_map[2] == source_paths
for path in source_paths:
    path.unlink(missing_ok=True)

fallback_visual_specs = [
    SlideSpec(title="Cover", visual="layout_id: cover"),
    SlideSpec(title="Why the old approach is reaching its limits", visual="layout_id: motivation_compare"),
    SlideSpec(title="Editable method", visual="layout_id: method_pipeline"),
    SlideSpec(title="Key adoption challenges", visual="layout_id: challenge_map"),
    SlideSpec(title="Summary", visual="layout_id: summary_takeaways"),
]
assert set(_selected_image_slide_indices(fallback_visual_specs, 2, include_cover=False)) == {2, 4}

non_image_layout_specs = [
    SlideSpec(title="Cover", visual="layout_id: cover"),
    SlideSpec(title="Editable process", visual="layout_id: method_pipeline"),
    SlideSpec(title="Editable metrics", visual="layout_id: benchmark_metrics"),
    SlideSpec(title="Summary", visual="layout_id: summary_takeaways"),
]
assert _selected_image_slide_indices(non_image_layout_specs, 1, include_cover=False) == [3]

fallback_source = TMP / "attachment-visual.png"
Image.new("RGB", (320, 180), (120, 45, 35)).save(fallback_source)
fallback_source_map = _map_source_visuals_to_slides(non_image_layout_specs, [fallback_source])
assert fallback_source_map == {3: [fallback_source]}, "an available attachment visual must be placed even without an explicit figure marker"
fallback_source.unlink(missing_ok=True)

empty_image_deck = Presentation()
empty_image_deck.slides.add_slide(empty_image_deck.slide_layouts[6])
empty_image_path = TMP / "minimum-image-empty.pptx"
empty_image_deck.save(empty_image_path)
minimum_image_issues = _ppt_image_binding_issues(empty_image_path, {})
assert minimum_image_issues and "至少需要 1 张" in minimum_image_issues[0]
empty_image_path.unlink(missing_ok=True)

hitsz_template = resolve_ppt_template(ASSET_ROOT, "hitsz")
scut_template = resolve_ppt_template(ASSET_ROOT, "scut")
assert hitsz_template.template_id == "hitsz" and scut_template.template_id == "scut"
assert hitsz_template.path and scut_template.path and hitsz_template.path != scut_template.path

sparse_specs = [
    SlideSpec(title="Cover", visual="layout_id: cover"),
    SlideSpec(
        title="Three supported takeaways",
        message="Evidence A • Evidence B • Evidence C",
        visual="layout_id: challenge_map",
        content_spec={"points": ["Evidence A", "Evidence B", "Evidence C"]},
    ),
]
sparse_repaired = _adapt_sparse_specs_for_quality_repair(sparse_specs, slide_numbers={2})
assert ppt_service._spec_layout_id(sparse_repaired[1]) == "summary_takeaways"
assert len(sparse_repaired[1].content_spec["cards"]) == 3


def windows_preview_contract_checks() -> None:
    source = TMP / "windows-preview-contract.pptx"
    target = TMP / "windows-preview-contract.pdf"
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_bytes(b"pptx-placeholder")

    def fake_run(command, **_kwargs):
        assert command[0] == sys.executable
        assert command[-1] == str(target)
        target.write_bytes(b"%PDF-preview")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    with patch.object(office_conversion.os, "name", "nt"), patch.object(
        office_conversion.subprocess,
        "run",
        side_effect=fake_run,
    ):
        pdf, error = office_conversion._convert_powerpoint_to_pdf_windows(TMP, source, target)
    assert pdf == target and error is None
    assert "KWPP.Application" in Path(office_conversion.__file__).read_text(encoding="utf-8")
    source.unlink(missing_ok=True)
    target.unlink(missing_ok=True)


windows_preview_contract_checks()

print("ppt slide-library regression passed")
