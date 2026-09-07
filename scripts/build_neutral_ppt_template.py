from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


GENERAL = {
    "bg": RGBColor(248, 250, 252),
    "ink": RGBColor(17, 24, 39),
    "muted": RGBColor(75, 85, 99),
    "accent": RGBColor(37, 99, 235),
    "accent2": RGBColor(15, 118, 110),
    "soft": RGBColor(239, 246, 255),
    "line": RGBColor(191, 208, 218),
}

SCHOOL_ACCENTS = {"C41230", "B61918", "FE2040", "FF0000"}


def remove_shape(shape) -> None:
    element = shape.element
    element.getparent().remove(element)


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def replace_school_color(color_format, replacement: RGBColor) -> None:
    try:
        rgb = color_format.rgb
        if rgb is not None and str(rgb).upper() in SCHOOL_ACCENTS:
            color_format.rgb = replacement
    except Exception:
        pass


def neutralize_shape_colors(shapes) -> None:
    for shape in walk_shapes(shapes):
        try:
            replace_school_color(shape.fill.fore_color, GENERAL["accent"])
        except Exception:
            pass
        try:
            replace_school_color(shape.line.color, GENERAL["accent"])
        except Exception:
            pass
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    try:
                        replace_school_color(run.font.color, GENERAL["accent"])
                    except Exception:
                        pass


def add_rect(slide, x: float, y: float, width: float, height: float, color: RGBColor, *, rounded: bool = False):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(slide, text: str, x: float, y: float, width: float, height: float, *, size: float, color: RGBColor, bold: bool = False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Comic Sans MS"
    properties = run._r.get_or_add_rPr()
    for tag, typeface in (("a:latin", "Comic Sans MS"), ("a:ea", "Microsoft YaHei"), ("a:cs", "Comic Sans MS")):
        element = properties.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            properties.append(element)
        element.set("typeface", typeface)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def build_neutral_template(source: Path, output: Path) -> None:
    prs = Presentation(str(source))

    for master in prs.slide_masters:
        master.background.fill.solid()
        master.background.fill.fore_color.rgb = GENERAL["bg"]
        for shape in list(master.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                remove_shape(shape)
        neutralize_shape_colors(master.shapes)
        for shape in master.shapes:
            if getattr(shape, "is_placeholder", False) and "TITLE" in str(shape.placeholder_format.type).upper():
                shape.fill.background()
                shape.line.fill.background()
        for layout in master.slide_layouts:
            for shape in list(layout.shapes):
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    remove_shape(shape)
            neutralize_shape_colors(layout.shapes)
            for shape in layout.shapes:
                if getattr(shape, "is_placeholder", False) and "TITLE" in str(shape.placeholder_format.type).upper():
                    shape.fill.background()
                    shape.line.fill.background()

    cover = prs.slides[0]
    for shape in list(cover.shapes):
        remove_shape(shape)
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = GENERAL["bg"]
    add_rect(cover, 0, 0, 13.333, 7.5, GENERAL["bg"])
    add_rect(cover, 0, 0, 13.333, 0.16, GENERAL["accent"])
    add_rect(cover, 0, 0.16, 4.4, 0.06, GENERAL["accent2"])
    add_rect(cover, 9.65, 1.2, 2.55, 4.95, GENERAL["soft"], rounded=True)
    add_rect(cover, 10.05, 1.65, 0.16, 3.95, GENERAL["accent2"], rounded=True)
    add_rect(cover, 10.52, 2.0, 0.16, 3.25, GENERAL["accent"], rounded=True)
    add_rect(cover, 10.99, 2.42, 0.16, 2.42, GENERAL["line"], rounded=True)
    kicker = add_text(cover, "GENERAL PRESENTATION", 1.18, 1.42, 4.6, 0.42, size=13, color=GENERAL["accent2"], bold=True)
    kicker.name = "neutral-cover-kicker"
    add_rect(cover, 1.18, 5.65, 7.55, 0.03, GENERAL["line"])
    add_text(cover, "Editable · structured · reusable", 1.18, 5.82, 5.6, 0.35, size=10, color=GENERAL["muted"])

    for slide in list(prs.slides)[1:]:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = GENERAL["bg"]
        neutralize_shape_colors(slide.shapes)
        for shape in slide.shapes:
            if str(getattr(shape, "name", "")).startswith("tpl-title-"):
                shape.fill.background()
                shape.line.fill.background()

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the neutral Janus multi-function PPT page library.")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_neutral_template(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
