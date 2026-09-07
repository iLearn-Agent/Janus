from __future__ import annotations

from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_NAMES = (
    "通用多功能PPT模板.pptx",
    "哈工深多功能PPT模板.pptx",
    "华工多功能PPT模板.pptx",
)
TEMPLATE_DIRS = (
    ROOT / "assets" / "departments" / "ppt_department" / "templates",
    ROOT / "departments" / "ppt_department" / "templates",
)

LATIN_FRIENDLY = "Comic Sans MS"
LATIN_FORMAL = "Arial"
CJK_FONT = "Microsoft YaHei"

FORMAL_PREFIXES = (
    "bench-kpi",
    "bench-table-",
    "big-kpi",
    "bar",
    "leaderboard-grid-",
    "ablation-grid-",
    "pipeline-grid-",
    "target-grid-",
    "wp-grid-",
    "eval-kpi",
    "eval-bar",
    "eval-status-grid-",
    "risk-grid-",
)

CARD_PREFIXES = (
    "basic-text-card",
    "motivation-current",
    "motivation-target",
    "motivation-gap",
    "challenge-card",
    "challenge-core",
    "pipeline-input",
    "pipeline-output",
    "pipeline-step-",
    "loop-node-",
    "loop-center",
    "loop-guardrail",
    "bench-kpi",
    "bench-protocol",
    "big-kpi",
    "big-interpretation",
    "bars-note",
    "leader-note",
    "ablation-callout",
    "ablation-guide",
    "summary-card",
    "summary-next",
    "target-root",
    "target-node",
    "domain-layers",
    "domain-node",
    "domain-legend",
    "route-stage-",
    "route-task-",
    "route-dependency",
    "wp-note",
    "risk-rule",
    "eval-kpi",
    "eval-scope",
    "eval-bars-panel",
    "roadmap-card-",
    "roadmap-note",
)

CAPTION_TOKENS = ("caption", "note", "guide", "legend", "rule")
FIGURE_CAPTION_PREFIXES = ("basic-caption", "evidence-caption-", "case-caption-main", "case-thumb-caption-")
LIGHT_LINE = RGBColor(205, 215, 225)
SEMANTIC_DARK = RGBColor(15, 58, 86)
SEMANTIC_PRIMARY = RGBColor(0, 138, 154)
SEMANTIC_PRIMARY_LIGHT = RGBColor(239, 248, 249)
SEMANTIC_MUTED = RGBColor(143, 176, 195)
SEMANTIC_MUTED_LIGHT = RGBColor(247, 250, 252)
SEMANTIC_WARM = RGBColor(224, 135, 55)
SEMANTIC_WARM_LIGHT = RGBColor(255, 247, 228)


def walk_shapes(shapes: Iterable):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def set_typefaces(run, latin: str) -> None:
    run.font.name = latin
    properties = run._r.get_or_add_rPr()
    for tag, typeface in (("a:latin", latin), ("a:ea", CJK_FONT), ("a:cs", latin)):
        element = properties.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            properties.append(element)
        element.set("typeface", typeface)


def formal_shape(name: str) -> bool:
    return name.startswith(FORMAL_PREFIXES)


def target_size(name: str, current: float | None) -> float | None:
    if name.startswith("tpl-title-"):
        return 28.0
    if current is None:
        return None
    if name in {"wp-note", "risk-rule"}:
        return max(current, 13.0)
    if formal_shape(name):
        return max(current, 11.0)
    if any(token in name for token in CAPTION_TOKENS):
        return max(current, 11.0)
    if name.startswith(CARD_PREFIXES):
        return max(current, 15.0)
    return max(current, 14.0)


def clear_effects(shape) -> None:
    try:
        properties = shape._element.spPr
        for tag in ("a:effectLst", "a:effectDag"):
            effect = properties.find(qn(tag))
            if effect is not None:
                properties.remove(effect)
    except Exception:
        pass


def style_card(shape) -> None:
    try:
        shape.line.color.rgb = LIGHT_LINE
        shape.line.dash_style = MSO_LINE_DASH_STYLE.SOLID
        shape.line.width = Pt(0.6)
    except Exception:
        pass
    clear_effects(shape)
    if getattr(shape, "has_text_frame", False):
        frame = shape.text_frame
        frame.margin_left = max(frame.margin_left, Inches(0.14))
        frame.margin_right = max(frame.margin_right, Inches(0.14))
        frame.margin_top = max(frame.margin_top, Inches(0.10))
        frame.margin_bottom = max(frame.margin_bottom, Inches(0.10))


def style_figure_caption(shape) -> None:
    try:
        shape.fill.background()
        shape.line.fill.background()
    except Exception:
        pass
    if getattr(shape, "has_text_frame", False):
        shape.text_frame.margin_left = Inches(0.04)
        shape.text_frame.margin_right = Inches(0.04)
        shape.text_frame.margin_top = Inches(0.02)
        shape.text_frame.margin_bottom = Inches(0.02)


def set_geometry(shape, x: float, y: float, width: float, height: float) -> None:
    if shape is None:
        return
    shape.left = Inches(x)
    shape.top = Inches(y)
    shape.width = Inches(width)
    shape.height = Inches(height)


def set_shape_colors(shape, *, fill: RGBColor | None = None, line: RGBColor | None = None) -> None:
    if shape is None:
        return
    try:
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = line
            shape.line.width = Pt(0.8)
    except Exception:
        pass


def named(slide, name: str):
    return next((shape for shape in walk_shapes(slide.shapes) if str(shape.name) == name), None)


def named_all(slide, name: str) -> list:
    return [shape for shape in walk_shapes(slide.shapes) if str(shape.name) == name]


def rectangle_edge_point(source, target) -> tuple[float, float]:
    source_x = float(source.left + source.width / 2)
    source_y = float(source.top + source.height / 2)
    target_x = float(target.left + target.width / 2)
    target_y = float(target.top + target.height / 2)
    delta_x = target_x - source_x
    delta_y = target_y - source_y
    if not delta_x and not delta_y:
        return source_x, source_y
    scale_x = float(source.width / 2) / abs(delta_x) if delta_x else float("inf")
    scale_y = float(source.height / 2) / abs(delta_y) if delta_y else float("inf")
    scale = min(scale_x, scale_y)
    return source_x + delta_x * scale, source_y + delta_y * scale


def set_line_endpoints(line, start: tuple[float, float], end: tuple[float, float]) -> None:
    if line is None:
        return
    start_x, start_y = start
    end_x, end_y = end
    line.left = int(min(start_x, end_x))
    line.top = int(min(start_y, end_y))
    line.width = max(1, int(abs(end_x - start_x)))
    line.height = max(1, int(abs(end_y - start_y)))
    transform = line._element.spPr.xfrm
    for attribute, enabled in (("flipH", end_x < start_x), ("flipV", end_y < start_y)):
        if enabled:
            transform.set(attribute, "1")
        else:
            transform.attrib.pop(attribute, None)


def remove_named(slide, name: str) -> None:
    for shape in list(walk_shapes(slide.shapes)):
        if str(getattr(shape, "name", "") or "") == name:
            element = shape._element
            if element.getparent() is not None:
                element.getparent().remove(element)


def style_minimal_grid(slide, stem: str, *, columns: int, rows: int, highlight_row: int | None = None) -> None:
    for row_index in range(1, rows + 1):
        if row_index == 1:
            fill = RGBColor(18, 61, 82)
        elif highlight_row is not None and row_index == highlight_row:
            fill = RGBColor(255, 244, 194)
        elif row_index % 2 == 0:
            fill = RGBColor(247, 250, 252)
        else:
            fill = RGBColor(255, 255, 255)
        for col in range(1, columns + 1):
            shape = named(slide, f"{stem}-r{row_index}c{col}")
            if shape is None:
                continue
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
            shape.line.fill.background()


def optimize_geometry(slide) -> None:
    role = "cover"
    for shape in slide.shapes:
        if str(shape.name).startswith("tpl-title-"):
            role = str(getattr(shape, "text", "") or "").strip()
            break

    if role == "basic_content":
        set_geometry(named(slide, "basic-text-card"), 0.75, 1.45, 3.38, 5.10)
        set_geometry(named(slide, "basic-image"), 4.43, 1.45, 8.15, 4.25)
        set_geometry(named(slide, "basic-image-cross-a"), 4.56, 1.58, 7.89, 3.99)
        set_geometry(named(slide, "basic-image-cross-b"), 4.56, 1.58, 7.89, 3.99)
        set_geometry(named(slide, "basic-caption"), 4.43, 5.84, 8.15, 0.72)
    elif role == "motivation_compare":
        current = named(slide, "motivation-current")
        target = named(slide, "motivation-target")
        gap = named(slide, "motivation-gap")
        arrow = named(slide, "motivation-arrow")
        set_geometry(current, 0.75, 1.72, 4.10, 3.45)
        set_geometry(target, 8.48, 1.72, 4.10, 3.45)
        set_geometry(gap, 5.24, 2.30, 2.85, 1.85)
        if current is not None and target is not None:
            center_y = float(current.top + current.height / 2)
            set_line_endpoints(arrow, (float(current.left + current.width), center_y), (float(target.left), center_y))
        set_shape_colors(current, fill=SEMANTIC_MUTED_LIGHT, line=RGBColor(218, 227, 233))
        set_shape_colors(target, fill=SEMANTIC_PRIMARY_LIGHT, line=RGBColor(192, 221, 222))
        set_shape_colors(gap, fill=SEMANTIC_WARM_LIGHT, line=SEMANTIC_WARM)
        for col in range(1, 5):
            set_geometry(named(slide, f"motivation-mini-table-r1c{col}"), 0.75 + (col - 1) * 2.91, 5.55, 2.91, 0.45)
            set_geometry(named(slide, f"motivation-mini-table-r2c{col}"), 0.75 + (col - 1) * 2.91, 6.00, 2.91, 0.56)
        style_minimal_grid(slide, "motivation-mini-table", columns=4, rows=2)
    elif role == "challenge_map":
        cards = [shape for shape in walk_shapes(slide.shapes) if str(shape.name) == "challenge-card"]
        positions = ((0.75, 1.50), (8.08, 1.50), (0.75, 4.82), (8.08, 4.82))
        for shape, (x, y) in zip(cards[:4], positions):
            set_geometry(shape, x, y, 4.50, 1.68)
        set_geometry(named(slide, "challenge-core"), 4.77, 2.93, 3.79, 1.62)
        set_geometry(named(slide, "challenge-note"), 4.12, 4.67, 5.09, 0.35)
    elif role == "method_pipeline":
        set_geometry(named(slide, "pipeline-input"), 0.73, 2.55, 1.82, 1.52)
        set_geometry(named(slide, "pipeline-input-cross-a"), 0.86, 2.68, 1.56, 1.26)
        set_geometry(named(slide, "pipeline-input-cross-b"), 0.86, 2.68, 1.56, 1.26)
        set_geometry(named(slide, "pipeline-output"), 10.78, 2.55, 1.82, 1.52)
        set_geometry(named(slide, "pipeline-output-cross-a"), 10.91, 2.68, 1.56, 1.26)
        set_geometry(named(slide, "pipeline-output-cross-b"), 10.91, 2.68, 1.56, 1.26)
        for index in range(1, 6):
            set_geometry(named(slide, f"pipeline-step-{index}"), 2.85 + (index - 1) * 1.526, 1.68, 1.33, 3.18)
            if index < 5:
                set_geometry(named(slide, f"pipeline-arrow-{index}"), 4.22 + (index - 1) * 1.526, 3.26, 0.16, 0.02)
            set_geometry(named(slide, f"pipeline-grid-r1c{index}"), 2.85 + (index - 1) * 1.526, 5.18, 1.526, 0.55)
            set_geometry(named(slide, f"pipeline-grid-r2c{index}"), 2.85 + (index - 1) * 1.526, 5.73, 1.526, 0.75)
    elif role == "method_loop":
        center = named(slide, "loop-center")
        set_geometry(center, 4.86, 2.79, 3.61, 1.52)
        positions = ((5.37, 1.43), (9.42, 3.02), (5.37, 4.75), (1.32, 3.02))
        nodes = [named(slide, f"loop-node-{name}") for name in ("Plan", "Act", "Reflect", "Observe")]
        for node, (x, y) in zip(nodes, positions):
            set_geometry(node, x, y, 2.59, 1.00)
        for index in range(4):
            set_line_endpoints(
                named(slide, f"loop-a{index + 1}"),
                rectangle_edge_point(nodes[index], nodes[(index + 1) % 4]),
                rectangle_edge_point(nodes[(index + 1) % 4], nodes[index]),
            )
        set_geometry(named(slide, "loop-guardrail"), 0.85, 5.97, 11.63, 0.62)
    elif role == "benchmark_metrics":
        for index in range(1, 4):
            x = 0.78 + (index - 1) * 2.66
            set_geometry(named(slide, f"bench-kpi{index}"), x, 1.48, 2.44, 2.08)
            set_geometry(named(slide, f"bench-kpi{index}-accent"), x, 1.48, 0.07, 2.08)
            set_geometry(named(slide, f"bench-kpi{index}-value"), x + 0.19, 1.68, 2.03, 0.66)
            set_geometry(named(slide, f"bench-kpi{index}-label"), x + 0.19, 2.45, 2.03, 0.42)
            set_geometry(named(slide, f"bench-kpi{index}-note"), x + 0.19, 2.95, 2.03, 0.42)
        set_geometry(named(slide, "bench-protocol"), 8.95, 1.48, 3.63, 2.08)
    elif role == "result_big_numbers":
        positions = {
            1: (0.85, 1.48, 5.28, 3.92, False),
            2: (6.45, 1.48, 6.03, 1.80, True),
            3: (6.45, 3.60, 6.03, 1.80, True),
        }
        for index, (x, y, width, height, compact) in positions.items():
            base = named(slide, f"big-kpi{index}")
            accent = named(slide, f"big-kpi{index}-accent")
            set_geometry(base, x, y, width, height)
            set_geometry(accent, x, y, 0.08, height)
            if compact:
                set_geometry(named(slide, f"big-kpi{index}-value"), x + 0.25, y + 0.22, 1.72, 0.60)
                set_geometry(named(slide, f"big-kpi{index}-label"), x + 2.08, y + 0.23, width - 2.34, 0.58)
                set_geometry(named(slide, f"big-kpi{index}-note"), x + 0.25, y + 1.02, width - 0.52, 0.52)
            else:
                set_geometry(named(slide, f"big-kpi{index}-value"), x + 0.28, y + 0.34, width - 0.56, 0.94)
                set_geometry(named(slide, f"big-kpi{index}-label"), x + 0.28, y + 1.50, width - 0.56, 0.55)
                set_geometry(named(slide, f"big-kpi{index}-note"), x + 0.28, y + 2.25, width - 0.56, 1.08)
            set_shape_colors(base, fill=SEMANTIC_PRIMARY_LIGHT if index == 1 else RGBColor(255, 255, 255), line=RGBColor(207, 221, 228))
            set_shape_colors(accent, fill=SEMANTIC_PRIMARY if index == 1 else SEMANTIC_MUTED, line=None)
            for suffix in ("-value", "-label", "-note"):
                set_shape_colors(named(slide, f"big-kpi{index}{suffix}"), fill=None, line=None)
        set_geometry(named(slide, "big-interpretation"), 0.85, 5.70, 11.63, 0.88)
        set_shape_colors(named(slide, "big-interpretation"), fill=SEMANTIC_WARM_LIGHT, line=None)
    elif role == "results_bars":
        set_geometry(named(slide, "bars-chart-panel"), 0.75, 1.48, 8.62, 5.10)
        set_geometry(named(slide, "bars-title"), 1.12, 1.80, 4.20, 0.42)
        for index in range(1, 5):
            y = 2.48 + (index - 1) * 0.92
            primary = index == 3
            bar_height = 0.46 if primary else 0.30
            set_geometry(named(slide, f"bar{index}-label"), 1.12, y, 1.45, 0.36)
            set_geometry(named(slide, f"bar{index}-track"), 2.82, y + (0.36 - bar_height) / 2, 4.75, bar_height)
            set_geometry(named(slide, f"bar{index}-fill"), 2.82, y + (0.36 - bar_height) / 2, 4.75, bar_height)
            set_geometry(named(slide, f"bar{index}-value"), 7.76, y, 0.76, 0.36)
            set_shape_colors(named(slide, f"bar{index}-track"), fill=RGBColor(233, 241, 245), line=None)
            set_shape_colors(named(slide, f"bar{index}-fill"), fill=SEMANTIC_PRIMARY if primary else SEMANTIC_MUTED, line=None)
        set_geometry(named(slide, "bars-note1"), 9.72, 1.48, 2.86, 5.10)
        set_shape_colors(named(slide, "bars-note1"), fill=SEMANTIC_PRIMARY_LIGHT, line=None)
        remove_named(slide, "bars-note2")
        remove_named(slide, "bars-note3")
    elif role == "leaderboard_table":
        widths = (0.95, 2.25, 1.45, 1.35, 2.55)
        for row_index in range(1, 8):
            x = 0.75
            for col, width in enumerate(widths, start=1):
                set_geometry(named(slide, f"leaderboard-grid-r{row_index}c{col}"), x, 1.52 + (row_index - 1) * 0.70, width, 0.70)
                x += width
        set_geometry(named(slide, "leader-note1"), 9.85, 1.52, 2.73, 5.02)
        remove_named(slide, "leader-note2")
        remove_named(slide, "leader-note3")
    elif role == "ablation_matrix":
        width = 8.25 / 6
        for row_index in range(1, 7):
            for col in range(1, 7):
                set_geometry(named(slide, f"ablation-grid-r{row_index}c{col}"), 0.75 + (col - 1) * width, 1.48 + (row_index - 1) * 0.82, width, 0.82)
        style_minimal_grid(slide, "ablation-grid", columns=6, rows=6, highlight_row=2)
        set_geometry(named(slide, "ablation-callout"), 9.35, 1.48, 3.23, 2.32)
        set_geometry(named(slide, "ablation-guide"), 9.35, 4.10, 3.23, 2.36)
    elif role == "evidence_grid":
        xs = (0.75, 4.57, 8.39)
        for row, y in enumerate((1.38, 4.04)):
            for col, x in enumerate(xs):
                slot = f"evidence-img-{row}-{col}"
                set_geometry(named(slide, slot), x, y, 3.57, 1.78)
                set_geometry(named(slide, f"{slot}-cross-a"), x + 0.11, y + 0.13, 3.35, 1.52)
                set_geometry(named(slide, f"{slot}-cross-b"), x + 0.11, y + 0.13, 3.35, 1.52)
                set_geometry(named(slide, f"evidence-caption-{row}-{col}"), x, y + 1.88, 3.57, 0.57)
    elif role == "case_gallery":
        set_geometry(named(slide, "case-main"), 0.75, 1.43, 7.18, 4.25)
        set_geometry(named(slide, "case-main-cross-a"), 0.88, 1.56, 6.92, 3.99)
        set_geometry(named(slide, "case-main-cross-b"), 0.88, 1.56, 6.92, 3.99)
        set_geometry(named(slide, "case-caption-main"), 0.75, 5.79, 7.18, 0.72)
        positions = ((8.23, 1.43), (10.40, 1.43), (8.23, 4.08), (10.40, 4.08))
        for index, (x, y) in enumerate(positions):
            set_geometry(named(slide, f"case-thumb-{index}"), x, y, 1.90, 1.55)
            set_geometry(named(slide, f"case-thumb-{index}-cross-a"), x + 0.10, y + 0.12, 1.70, 1.31)
            set_geometry(named(slide, f"case-thumb-{index}-cross-b"), x + 0.10, y + 0.12, 1.70, 1.31)
            set_geometry(named(slide, f"case-thumb-caption-{index}"), x, y + 1.66, 1.90, 0.72)
    elif role == "summary_takeaways":
        set_geometry(named(slide, "summary-card1"), 0.85, 1.58, 5.42, 3.85)
        set_geometry(named(slide, "summary-card2"), 6.67, 1.58, 5.81, 1.67)
        set_geometry(named(slide, "summary-card3"), 6.67, 3.76, 5.81, 1.67)
        set_geometry(named(slide, "summary-next"), 0.85, 5.73, 11.63, 0.85)
        for index in range(1, 4):
            shape = named(slide, f"summary-card{index}")
            if shape is not None:
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_shape_colors(named(slide, "summary-card1"), fill=SEMANTIC_PRIMARY_LIGHT, line=None)
        set_shape_colors(named(slide, "summary-card2"), fill=RGBColor(255, 255, 255), line=RGBColor(218, 227, 233))
        set_shape_colors(named(slide, "summary-card3"), fill=RGBColor(255, 255, 255), line=RGBColor(218, 227, 233))
        set_shape_colors(named(slide, "summary-next"), fill=SEMANTIC_DARK, line=None)
    elif role == "project_target_map":
        root = named(slide, "target-root")
        set_geometry(root, 0.85, 1.48, 11.63, 1.08)
        nodes = named_all(slide, "target-node")
        links = named_all(slide, "target-link")
        width, gap, start = 2.76, 0.20, 0.86
        for index, node in enumerate(nodes[:4]):
            set_geometry(node, start + index * (width + gap), 3.10, width, 1.66)
            set_line_endpoints(
                links[index],
                rectangle_edge_point(root, node),
                rectangle_edge_point(node, root),
            )
        strip_width = 11.63 / 4
        for col in range(1, 5):
            set_geometry(named(slide, f"target-grid-r1c{col}"), 0.85 + (col - 1) * strip_width, 5.61, strip_width, 0.42)
            set_geometry(named(slide, f"target-grid-r2c{col}"), 0.85 + (col - 1) * strip_width, 6.03, strip_width, 0.55)
        style_minimal_grid(slide, "target-grid", columns=4, rows=2)
    elif role == "domain_object_map":
        layer_shape = named(slide, "domain-layers")
        set_geometry(layer_shape, 0.75, 1.48, 2.35, 4.97)
        if layer_shape is not None:
            layer_shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
            layer_shape.text_frame.margin_left = Inches(0.26)
            layer_shape.text_frame.margin_right = Inches(0.20)
            layer_shape.text_frame.margin_top = Inches(0.28)
        positions = ((3.48, 1.78), (6.54, 1.78), (9.60, 1.78), (5.01, 4.00), (8.07, 4.00))
        nodes = named_all(slide, "domain-node")
        for node, (x, y) in zip(nodes[:5], positions):
            set_geometry(node, x, y, 2.45, 1.20)
        pairs = ((0, 1), (1, 2), (0, 3), (1, 3), (1, 4), (2, 4))
        for edge, (source_index, target_index) in zip(named_all(slide, "domain-edge"), pairs):
            source, target = nodes[source_index], nodes[target_index]
            set_line_endpoints(edge, rectangle_edge_point(source, target), rectangle_edge_point(target, source))
        set_geometry(named(slide, "domain-legend"), 3.48, 5.75, 8.57, 0.70)
    elif role == "technical_route":
        left, right, gap, count = 0.75, 12.58, 0.28, 4
        width = (right - left - gap * (count - 1)) / count
        headers = []
        for index in range(count):
            x = left + index * (width + gap)
            header = named(slide, f"route-stage-{index}")
            task = named(slide, f"route-task-{index}")
            headers.append(header)
            set_geometry(header, x, 1.50, width, 0.78)
            set_geometry(task, x, 2.56, width, 2.95)
            if task is not None:
                task.text_frame.vertical_anchor = MSO_ANCHOR.TOP
                task.text_frame.margin_top = Inches(0.24)
                task.text_frame.margin_left = Inches(0.20)
                task.text_frame.margin_right = Inches(0.18)
        for index in range(3):
            set_line_endpoints(
                named(slide, f"route-arrow-{index}"),
                rectangle_edge_point(headers[index], headers[index + 1]),
                rectangle_edge_point(headers[index + 1], headers[index]),
            )
        set_geometry(named(slide, "route-dependency"), 0.75, 5.91, 11.83, 0.67)
    elif role == "evaluation_dashboard":
        for index in range(1, 4):
            x = 0.81 + (index - 1) * 2.58
            set_geometry(named(slide, f"eval-kpi{index}"), x, 1.48, 2.31, 2.12)
            set_geometry(named(slide, f"eval-kpi{index}-accent"), x, 1.48, 0.07, 2.12)
            set_geometry(named(slide, f"eval-kpi{index}-value"), x + 0.19, 1.68, 1.93, 0.62)
            set_geometry(named(slide, f"eval-kpi{index}-label"), x + 0.20, 2.36, 1.90, 0.56)
            set_geometry(named(slide, f"eval-kpi{index}-note"), x + 0.20, 3.00, 1.90, 0.44)
            set_shape_colors(named(slide, f"eval-kpi{index}"), fill=SEMANTIC_PRIMARY_LIGHT if index == 1 else RGBColor(255, 255, 255), line=RGBColor(218, 227, 233))
            set_shape_colors(named(slide, f"eval-kpi{index}-accent"), fill=SEMANTIC_PRIMARY if index == 1 else SEMANTIC_MUTED, line=None)
            for suffix in ("-value", "-label", "-note"):
                set_shape_colors(named(slide, f"eval-kpi{index}{suffix}"), fill=None, line=None)
        set_geometry(named(slide, "eval-scope"), 8.78, 1.48, 3.80, 2.12)
        set_shape_colors(named(slide, "eval-scope"), fill=SEMANTIC_WARM_LIGHT, line=None)
    elif role == "workpackage_matrix":
        widths = (1.15, 1.35, 2.55, 2.45, 1.85, 1.80)
        for row_index in range(1, 7):
            x = 0.75
            for col, width in enumerate(widths, start=1):
                set_geometry(named(slide, f"wp-grid-r{row_index}c{col}"), x, 1.48 + (row_index - 1) * 0.70, width, 0.70)
                x += width
        style_minimal_grid(slide, "wp-grid", columns=6, rows=6, highlight_row=2)
        set_geometry(named(slide, "wp-note"), 0.75, 5.93, 11.15, 0.64)
    elif role == "risk_action_table":
        widths = (2.0, 2.2, 2.35, 3.75, 1.2)
        for row_index in range(1, 7):
            x = 0.75
            for col, width in enumerate(widths, start=1):
                set_geometry(named(slide, f"risk-grid-r{row_index}c{col}"), x, 1.48 + (row_index - 1) * 0.70, width, 0.70)
                x += width
        style_minimal_grid(slide, "risk-grid", columns=5, rows=6)
        set_geometry(named(slide, "risk-rule"), 0.75, 5.93, 11.50, 0.64)
    elif role == "milestone_roadmap":
        axis_y, card_width, card_height = 3.72, 1.85, 1.30
        left = 0.85 + card_width / 2
        right = 12.48 - card_width / 2
        centers = [left + index * (right - left) / 5 for index in range(6)]
        for index, center_x in enumerate(centers):
            card = named(slide, f"roadmap-card-{index}")
            dot = named(slide, f"roadmap-dot-{index}")
            link = named(slide, f"roadmap-link-{index}")
            card_y = 1.72 if index % 2 == 0 else 4.30
            set_geometry(card, center_x - card_width / 2, card_y, card_width, card_height)
            set_geometry(dot, center_x - 0.13, axis_y - 0.13, 0.26, 0.26)
            if index % 2 == 0:
                start = (float(card.left + card.width / 2), float(card.top + card.height))
                end = (float(dot.left + dot.width / 2), float(dot.top))
            else:
                start = (float(dot.left + dot.width / 2), float(dot.top + dot.height))
                end = (float(card.left + card.width / 2), float(card.top))
            set_line_endpoints(link, start, end)
        set_geometry(named(slide, "roadmap-axis"), centers[0], axis_y, centers[-1] - centers[0], 0.01)
        set_geometry(named(slide, "roadmap-note"), 0.85, 5.98, 11.63, 0.61)
    elif role == "media_showcase":
        set_geometry(named(slide, "media-showcase-main"), 0.60, 1.25, 12.13, 5.45)
        set_geometry(named(slide, "media-showcase-main-cross-a"), 0.73, 1.38, 11.87, 5.19)
        set_geometry(named(slide, "media-showcase-main-cross-b"), 0.73, 1.38, 11.87, 5.19)


def optimize_template(path: Path) -> None:
    prs = Presentation(str(path))
    shape_collections = []
    for master in prs.slide_masters:
        shape_collections.append(master.shapes)
        for layout in master.slide_layouts:
            shape_collections.append(layout.shapes)
    shape_collections.extend(slide.shapes for slide in prs.slides)

    for shapes in shape_collections:
        for shape in walk_shapes(shapes):
            name = str(getattr(shape, "name", "") or "")
            if name.startswith(CARD_PREFIXES):
                style_card(shape)
            if name.startswith(FIGURE_CAPTION_PREFIXES):
                style_figure_caption(shape)
            if not getattr(shape, "has_text_frame", False):
                continue
            latin = LATIN_FORMAL if formal_shape(name) else LATIN_FRIENDLY
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    set_typefaces(run, latin)
                    size = target_size(name, run.font.size.pt if run.font.size is not None else None)
                    if size is not None:
                        run.font.size = Pt(size)

    for slide in prs.slides:
        optimize_geometry(slide)
    prs.save(str(path))
    print(path.relative_to(ROOT))


def main() -> None:
    paths = [directory / name for directory in TEMPLATE_DIRS for name in TEMPLATE_NAMES]
    missing = [path for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing PPT templates: " + ", ".join(str(path) for path in missing))
    for path in paths:
        optimize_template(path)


if __name__ == "__main__":
    main()
