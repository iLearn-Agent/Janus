from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIRS = (
    ROOT / "assets" / "departments" / "ppt_department" / "templates",
    ROOT / "departments" / "ppt_department" / "templates",
)
TEMPLATES = tuple(
    template_dir / filename
    for template_dir in TEMPLATE_DIRS
    for filename in ("华工多功能PPT模板.pptx", "哈工深多功能PPT模板.pptx")
)

SEMANTIC_PREFIXES = (
    "basic-",
    "motivation-",
    "challenge-",
    "pipeline-",
    "loop-",
    "bench-",
    "big-",
    "bars-",
    "bar",
    "leader-",
    "ablation-",
    "evidence-",
    "case-",
    "summary-",
    "target-",
    "domain-",
    "route-",
    "wp-",
    "eval-",
    "risk-",
    "roadmap-",
    "media-",
)

TARGET_BOTTOM = 6.55
ALREADY_EXPANDED_BOTTOM = 6.45
MAX_SCALE = 1.45


def _is_semantic(shape) -> bool:
    return str(getattr(shape, "name", "") or "").startswith(SEMANTIC_PREFIXES)


def _shape_named(slide, name: str):
    return next((shape for shape in slide.shapes if str(shape.name) == name), None)


def _set_geometry(shape, x: float, y: float, width: float, height: float) -> None:
    if shape is None:
        return
    shape.left = Inches(x)
    shape.top = Inches(y)
    shape.width = Inches(width)
    shape.height = Inches(height)


def _remove_named(slide, name: str) -> None:
    for shape in list(slide.shapes):
        if str(shape.name) == name:
            shape._element.getparent().remove(shape._element)


def _layout_role(slide) -> str:
    for shape in slide.shapes:
        if shape.name.startswith("tpl-title-") and str(getattr(shape, "text", "") or "").strip():
            return str(shape.text or "").strip()
    return "cover"


def _apply_layout_fixes(slide) -> None:
    role = _layout_role(slide)
    if role == "basic_content":
        # The basic visual page should be image-dominant and should not expose
        # internal column labels such as TEXT/VISUAL or 内容/视觉.
        _remove_named(slide, "tpl-kicker-label")
        _set_geometry(_shape_named(slide, "basic-text-card"), 0.75, 1.45, 4.15, 5.1)
        _set_geometry(_shape_named(slide, "basic-image"), 5.15, 1.45, 7.43, 4.4)
        _set_geometry(_shape_named(slide, "basic-image-cross-a"), 5.28, 1.6, 7.17, 4.1)
        _set_geometry(_shape_named(slide, "basic-image-cross-b"), 5.28, 1.6, 7.17, 4.1)
        _set_geometry(_shape_named(slide, "basic-caption"), 5.15, 5.98, 7.43, 0.57)
    elif role == "motivation_compare":
        # Reserve a real center column for the transition claim. The previous
        # 1.8-inch callout forced even short sentences into unreadable text.
        _set_geometry(_shape_named(slide, "motivation-current"), 0.75, 1.55, 4.25, 3.72)
        _set_geometry(_shape_named(slide, "motivation-target"), 8.33, 1.55, 4.25, 3.72)
        _set_geometry(_shape_named(slide, "motivation-gap"), 5.18, 2.28, 2.97, 1.08)
        _set_geometry(_shape_named(slide, "motivation-arrow"), 5.38, 3.72, 2.57, 0.02)
        for name in ("motivation-current", "motivation-target"):
            shape = _shape_named(slide, name)
            if shape is not None and getattr(shape, "has_text_frame", False):
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def _expand_slide(slide) -> tuple[float, float, float]:
    shapes = [shape for shape in slide.shapes if _is_semantic(shape)]
    if not shapes:
        return 0.0, 0.0, 1.0

    top = min(float(shape.top) / Inches(1) for shape in shapes)
    bottom = max(float(shape.top + shape.height) / Inches(1) for shape in shapes)
    if bottom >= ALREADY_EXPANDED_BOTTOM or bottom <= top:
        return top, bottom, 1.0

    scale = min(MAX_SCALE, (TARGET_BOTTOM - top) / (bottom - top))
    for shape in shapes:
        old_top = float(shape.top) / Inches(1)
        old_height = float(shape.height) / Inches(1)
        shape.top = Inches(top + (old_top - top) * scale)
        shape.height = Inches(old_height * scale)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                row.height = int(row.height * scale)
    return top, bottom, scale


def expand_template(path: Path) -> None:
    prs = Presentation(str(path))
    changes: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        top, bottom, scale = _expand_slide(slide)
        _apply_layout_fixes(slide)
        if scale > 1.001:
            role = _layout_role(slide)
            changes.append(f"{index:02d} {role}: {bottom:.2f} -> {top + (bottom - top) * scale:.2f} ({scale:.2f}x)")
    prs.save(str(path))
    print(path.relative_to(ROOT))
    for change in changes:
        print(f"  {change}")


def main() -> None:
    for path in TEMPLATES:
        if not path.is_file():
            raise FileNotFoundError(path)
        expand_template(path)


if __name__ == "__main__":
    main()
